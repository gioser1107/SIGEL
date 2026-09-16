#!/usr/bin/env python3
"""SRS definitivo: formato de María + Guía Trayecto III V2, IDs originales."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from _generar_srs_uptaeb import RF_FICHAS, RF_MATRIZ, RNF_FICHAS, RNF_MATRIZ
from _subrequisitos_srs import SUB_RCI, aplicar_fichas_rci, aplicar_matriz

RAIZ = Path(__file__).resolve().parent
SALIDA = RAIZ / "ENTREGA_INGENIERIA_SOFTWARE" / "modelo_negocio" / "SRS_FORMATO_ESTANDAR_TRAYECTO_III.docx"

TITULO = (
    "DESARROLLO DE UN SISTEMA INTEGRAL PARA LA GESTIÓN LOGÍSTICA TURÍSTICA "
    "CON MÓDULOS DE ANALÍTICA PREDICTIVA Y ARQUITECTURA MVC PARA LA AGENCIA "
    "TRAVEL BQTO, PARROQUIA ANA SOTO, BARQUISIMETO"
)

RCI_MATRIZ = [
    (
        "RCI-01",
        "Predictor de rentabilidad de viajes",
        "Inteligente",
        "Alta",
        "Técnica: Regresión lineal múltiple (Python).\n"
        "Métrica de Desempeño: error absoluto medio del punto de equilibrio < 15 % sobre el conjunto de prueba.\n"
        "Origen de Datos: costos operativos, precio de venta, cupos vendidos y mes del viaje extraídos de SIGEL.",
        "Viajes que generan pérdidas por estimaciones manuales de costos antes de publicar.",
    ),
    (
        "RCI-02",
        "Motor de recomendación de destinos",
        "Inteligente",
        "Media",
        "Técnica: Filtrado colaborativo o K-Nearest Neighbors.\n"
        "Métrica de Desempeño: sugerencia de 3 destinos con índice de relevancia ≥ 75 %.\n"
        "Origen de Datos: reservas, destinos y reseñas almacenados en SIGEL.",
        "Baja personalización de la oferta; ATC no dispone de tiempo para recomendar destinos a clientes recurrentes.",
    ),
    (
        "RCI-03",
        "Predictor de demanda estacional",
        "Inteligente",
        "Alta",
        "Técnica: Series temporales (ARIMA o equivalente).\n"
        "Métrica de Desempeño: proyección a 30 días del volumen de asientos con error < 15 %.\n"
        "Origen de Datos: serie diaria de reservas y asientos vendidos por destino y temporada.",
        "Asignación ineficiente de transporte por desconocimiento de picos de demanda.",
    ),
]

RCI_MATRIZ = aplicar_matriz(RCI_MATRIZ, SUB_RCI)
MATRIZ = list(RF_MATRIZ) + list(RNF_MATRIZ) + RCI_MATRIZ

RCI_FICHAS = [
    (
        "RCI-01",
        "Predictor de rentabilidad de viajes",
        "Regresión lineal múltiple implementada en Python.",
        "El modelo toma costos de transporte, pago de guías, viáticos, precio de venta y ocupación histórica para estimar el punto de equilibrio del viaje. El resultado se expresa como ocupación mínima requerida y como indicador de ganancia o pérdida esperada antes de publicar.",
        "Histórico operativo de Travel BQTO extraído de SIGEL: costo de transporte, gastos del destino, precio de venta, puestos vendidos y mes del viaje.",
        "Impedir ofertar paquetes con rentabilidad nula o negativa y apoyar la decisión de publicar o no un viaje.",
    ),
    (
        "RCI-02",
        "Motor de recomendación de destinos",
        "Filtrado colaborativo o K-Nearest Neighbors.",
        "El algoritmo relaciona el historial de reservas y reseñas del cliente con destinos similares. Devuelve tres destinos ordenados por relevancia para apoyar a atención al cliente y al portal.",
        "Reservas, destinos y reseñas persistidos en SIGEL; no se usan datos de género ni edad porque el maestro de clientes no los almacena.",
        "Aumentar la personalización de la oferta y reducir el tiempo que ATC dedica a recomendar destinos de forma empírica.",
    ),
    (
        "RCI-03",
        "Predictor de demanda estacional",
        "Análisis de series temporales (ARIMA o equivalente).",
        "El modelo descompone la serie diaria de asientos vendidos para detectar estacionalidad (asuetos, vacaciones escolares) y proyecta la demanda a 30 días. El resultado dimensiona la unidad de transporte a contratar.",
        "Serie temporal de reservas y asientos vendidos por fecha y destino, con indicador de temporada alta o baja.",
        "Anticipar la contratación de unidades y evitar sobrecostos por contrataciones de emergencia en picos de demanda.",
    ),
]
RCI_FICHAS = aplicar_fichas_rci(RCI_FICHAS, SUB_RCI)

VALIDACION_RCI = [
    (
        "Costos operativos totales, precio del paquete, histórico de puestos vendidos del mes",
        "Ocupación mínima para equilibrio y probabilidad de ganancia o pérdida",
        "Blinda las finanzas de Travel BQTO al evitar la ejecución de viajes a pérdida.",
    ),
    (
        "Historial de reservas y reseñas del cliente, catálogo de destinos activos",
        "Tres destinos recomendados con índice de relevancia",
        "Mejora la captación de clientes recurrentes y reduce la carga de ATC.",
    ),
    (
        "Serie histórica de fechas, asientos vendidos por día, indicador de temporada o feriado",
        "Proyección del volumen de asientos a 30 días, error < 15 %",
        "Permite negociar transporte con antelación y evitar sobrecostos en picos de demanda.",
    ),
]


def _arial(run, size=12, bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _p(doc, texto, size=12, bold=False, center=False, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(texto)
    _arial(run, size, bold)
    return p


def _sombrear(celda, hex_color):
    tc = celda._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _celda(celda, texto, bold=False, size=12, fill=None):
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(texto)
    _arial(run, size, bold)
    if fill:
        _sombrear(celda, fill)
    for par in celda.paragraphs:
        par.paragraph_format.line_spacing = 1.08


def _merge(tabla, fila, c0=0, c1=1):
    tabla.rows[fila].cells[c0].merge(tabla.rows[fila].cells[c1])


def _setup(doc):
    for seccion in doc.sections:
        seccion.page_width = Cm(21.59)
        seccion.page_height = Cm(27.94)
        seccion.top_margin = Cm(2.5)
        seccion.bottom_margin = Cm(2.5)
        seccion.left_margin = Cm(2.5)
        seccion.right_margin = Cm(2.5)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(12)


def _tabla_matriz(doc):
    tabla = doc.add_table(rows=1, cols=6)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [
        "ID",
        "Nombre del Requisito",
        "Tipo",
        "Prioridad",
        "Especificación Técnica y Criterios de Aceptación",
        "Necesidad o Problema Asociado",
    ]
    anchos = [Cm(1.8), Cm(3.2), Cm(2.1), Cm(2.0), Cm(5.4), Cm(3.5)]
    for i, h in enumerate(headers):
        _celda(tabla.rows[0].cells[i], h, bold=True, size=10, fill="D9E2F3")
    for fila in MATRIZ:
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            _celda(celdas[i], valor, bold=(i == 0), size=9)
    for fila in tabla.rows:
        for i, celda in enumerate(fila.cells):
            celda.width = anchos[i]


def _ficha_rf(doc, ficha):
    ident, nombre, desc, actores, entradas, salidas, pre = ficha
    t = doc.add_table(rows=5, cols=2)
    t.style = "Table Grid"
    _celda(t.rows[0].cells[0], f"ID: {ident}", bold=True, size=12, fill="D9E2F3")
    _celda(t.rows[0].cells[1], f"Nombre: {nombre}", bold=True, size=12, fill="D9E2F3")
    _merge(t, 1)
    _celda(t.rows[1].cells[0], f"Descripción: {desc}", size=12)
    _merge(t, 2)
    _celda(t.rows[2].cells[0], f"Actores: {actores}", size=12)
    _celda(t.rows[3].cells[0], f"Entradas: {entradas}", size=12)
    _celda(t.rows[3].cells[1], f"Salidas: {salidas}", size=12)
    _merge(t, 4)
    _celda(t.rows[4].cells[0], f"Precondición: {pre}", size=12)
    doc.add_paragraph()


def _ficha_rnf(doc, ficha):
    ident, nombre, desc, metrica, resultados, espec, prioridad = ficha
    t = doc.add_table(rows=6, cols=2)
    t.style = "Table Grid"
    _celda(t.rows[0].cells[0], f"ID: {ident}", bold=True, size=12, fill="E2EFDA")
    _celda(t.rows[0].cells[1], f"Nombre: {nombre}", bold=True, size=12, fill="E2EFDA")
    for i, texto in enumerate(
        [
            f"Descripción: {desc}",
            f"Métrica y evaluación: {metrica}",
            f"Resultados: {resultados}",
            f"Especificación: {espec}",
            f"Prioridad: {prioridad}",
        ],
        start=1,
    ):
        _merge(t, i)
        _celda(t.rows[i].cells[0], texto, size=12)
    doc.add_paragraph()


def _ficha_rci(doc, ficha):
    ident, nombre, tecnica, descripcion, datos, objetivo = ficha
    t = doc.add_table(rows=5, cols=1)
    t.style = "Table Grid"
    _celda(t.rows[0].cells[0], f"ID: {ident} - {nombre}", bold=True, size=12, fill="FFF2CC")
    _celda(t.rows[1].cells[0], f"Técnica de IA: {tecnica}", size=12)
    _celda(t.rows[2].cells[0], f"Descripción del Algoritmo: {descripcion}", size=12)
    _celda(t.rows[3].cells[0], f"Datos de Entrenamiento: {datos}", size=12)
    _celda(t.rows[4].cells[0], f"Objetivo / Predicción: {objetivo}", size=12)
    doc.add_paragraph()


def generar():
    doc = Document()
    _setup(doc)

    _p(doc, "PROGRAMA NACIONAL DE FORMACIÓN EN INFORMÁTICA", 12, True, True, 10)
    _p(doc, TITULO, 12, True, True, 6)
    _p(doc, "PERIODO SEPTIEMBRE 2026", 12, True, True, 14)
    _p(doc, "Especificación de Requisitos de Software (SRS)", 14, True, True, 8)
    _p(doc, "Guía Estándar Trayecto III V2 — ISO/IEC/IEEE 29148", 12, False, True, 8)
    _p(doc, "Integrantes: Sergio Jiménez, Gabriel Jiménez", 12, False, True, 4)
    _p(doc, "Tutor: Edecio Freitez", 12, False, True, 4)
    _p(doc, "Sistema: SIGEL — Travel BQTO", 12, False, True, 16)

    _p(doc, "1. Nomenclatura y clasificación", 12, True, False, 8, 6)
    _p(
        doc,
        "RF-: Requisito Funcional (define qué hace el sistema). "
        "RNF-: Requisito No Funcional (define bajo qué calidad opera). "
        "RCI-: Requisito del Componente Inteligente (lógica avanzada e inteligencia artificial). "
        "Los identificadores RF-01 a RF-10 conservan la numeración del documento original del equipo.",
        12,
        False,
        False,
        10,
    )

    _p(doc, "2. Matriz Única de Especificación de Requisitos (SRS)", 12, True, False, 8)
    _p(
        doc,
        "Se conservan los requisitos funcionales RF-01 a RF-10 del documento original y se agregan RF-11 a RF-19 "
        "para cubrir clientes, reservas, portales, reportes, flota, catálogo financiero, viajes y bitácora. "
        "Los RNF-01 a RNF-10 son los del documento original, actualizados con métricas verificables; "
        "RNF-11 a RNF-18 cubren permisos, roles, módulos, bitácora y rendimiento exigidos por la guía. "
        "Los RCI-01 a RCI-03 se mantienen como componente inteligente de Trayecto III.",
        12,
        False,
        False,
        8,
    )
    _tabla_matriz(doc)

    doc.add_page_break()
    _p(doc, "3. Fichas Técnicas de Detalle", 12, True, False, 8)
    _p(doc, "A. Fichas para Requisitos Funcionales (RF)", 12, True, False, 10)
    for ficha in RF_FICHAS:
        _ficha_rf(doc, ficha)

    doc.add_page_break()
    _p(doc, "B. Fichas para Requisitos No Funcionales (RNF)", 12, True, False, 10)
    for ficha in RNF_FICHAS:
        _ficha_rnf(doc, ficha)

    doc.add_page_break()
    _p(doc, "C. Fichas para el Componente Inteligente (RCI)", 12, True, False, 10)
    _p(
        doc,
        "El RF-15 cubre la estadística descriptiva operativa. Los RCI especifican la capa predictiva "
        "exigida por Trayecto III, con el histórico que genera SIGEL (costos, reservas, destinos y reseñas).",
        12,
        False,
        False,
        10,
    )
    for ficha in RCI_FICHAS:
        _ficha_rci(doc, ficha)

    _p(doc, "4. Matriz de Validación del Componente Inteligente", 12, True, False, 10, 8)
    val = doc.add_table(rows=1, cols=3)
    val.style = "Table Grid"
    for i, h in enumerate(
        ["Variable de Entrada", "Variable de Salida (Resultado IA)", "Impacto en la Comunidad"]
    ):
        _celda(val.rows[0].cells[i], h, bold=True, size=12, fill="FFF2CC")
    for fila in VALIDACION_RCI:
        celdas = val.add_row().cells
        for i, valor in enumerate(fila):
            _celda(celdas[i], valor, size=12)

    return doc


if __name__ == "__main__":
    print("Catálogo RCI para el SRS oficial. Regenerar con: python3 _rellenar_srs_original.py")
