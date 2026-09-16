#!/usr/bin/env python3
"""Rellena el Word plantilla UPTAEB: mismo formato, información actualizada."""

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from _generar_srs_formato_uptaeb import RCI_FICHAS, RCI_MATRIZ, VALIDACION_RCI
from _generar_srs_uptaeb import RF_FICHAS, RF_MATRIZ, RNF_FICHAS, RNF_MATRIZ

ORIGEN = Path("/Users/sergiojimenez/Downloads/REQUERIMIENTOS.docx")
RAIZ = Path(__file__).resolve().parent
DESTINO = RAIZ / "ENTREGA_INGENIERIA_SOFTWARE" / "modelo_negocio" / "SRS_FORMATO_ESTANDAR_TRAYECTO_III.docx"
INTEGRANTES_SALIDOS = ("Mar", "ía", " Alvarado", "Luis Herice")


def _set_plain(celda, texto):
    texto = texto or ""
    pars = celda.paragraphs
    if not pars:
        return
    primero = pars[0]
    if primero.runs:
        primero.runs[0].text = texto
        for run in primero.runs[1:]:
            run.text = ""
    else:
        run = primero.add_run(texto)
        run.font.name = "Arial"
    for extra in pars[1:]:
        for run in extra.runs:
            run.text = ""


def _set_etiqueta(celda, etiqueta, valor):
    texto = f" {valor}" if valor else ""
    pars = celda.paragraphs
    if not pars:
        return
    p = pars[0]
    if len(p.runs) >= 2:
        p.runs[0].text = etiqueta
        p.runs[0].bold = True
        p.runs[1].text = texto
        for run in p.runs[2:]:
            run.text = ""
    elif p.runs:
        p.runs[0].text = f"{etiqueta}{texto}"
    else:
        r0 = p.add_run(etiqueta)
        r0.bold = True
        r0.font.name = "Arial"
        r1 = p.add_run(texto)
        r1.font.name = "Arial"
    for extra in pars[1:]:
        for run in extra.runs:
            run.text = ""


def _id_rnf(numero):
    return f"RNF - {numero:02d}"


def _rellenar_matriz(tabla):
    filas = (
        list(RF_MATRIZ)
        + [(_id_rnf(i + 1), *fila[1:]) for i, fila in enumerate(RNF_MATRIZ)]
        + list(RCI_MATRIZ)
    )
    while len(tabla.rows) - 1 < len(filas):
        src = tabla.rows[1]._tr
        tabla.rows[-1]._tr.addnext(deepcopy(src))

    for i, fila in enumerate(filas, start=1):
        celdas = tabla.rows[i].cells
        for j, valor in enumerate(fila):
            _set_plain(celdas[j], " ".join(valor.split()))


def _rellenar_rf(tabla, ficha):
    ident, nombre, desc, actores, entradas, salidas, pre = ficha
    _set_etiqueta(tabla.rows[0].cells[0], "ID:", ident)
    _set_etiqueta(tabla.rows[0].cells[1], "Nombre:", nombre)
    _set_etiqueta(tabla.rows[1].cells[0], "Descripción:", desc)
    _set_etiqueta(tabla.rows[2].cells[0], "Actores:", actores)
    _set_etiqueta(tabla.rows[3].cells[0], "Entradas:", entradas)
    _set_etiqueta(tabla.rows[3].cells[1], "Salidas:", salidas)
    _set_etiqueta(tabla.rows[4].cells[0], "Precondición:", pre)


def _rellenar_rnf(tabla, ficha):
    ident, nombre, desc, metrica, resultados, espec, prioridad = ficha
    _set_etiqueta(tabla.rows[0].cells[0], "ID:", ident)
    _set_etiqueta(tabla.rows[0].cells[1], "Nombre:", nombre)
    _set_etiqueta(tabla.rows[1].cells[0], "Descripción:", desc)
    _set_etiqueta(tabla.rows[2].cells[0], "Métrica y evaluación:", metrica)
    _set_etiqueta(tabla.rows[3].cells[0], "Resultados:", resultados)
    _set_etiqueta(tabla.rows[4].cells[0], "Especificación:", espec)
    _set_etiqueta(tabla.rows[5].cells[0], "Prioridad:", prioridad)


def _rellenar_rci(tabla, ficha):
    ident, nombre, tecnica, descripcion, datos, objetivo = ficha
    _set_plain(tabla.rows[0].cells[0], f"ID: {ident} - {nombre}")
    _set_plain(tabla.rows[1].cells[0], f"Técnica de IA: {tecnica}")
    _set_plain(tabla.rows[2].cells[0], f"Descripción del Algoritmo: {descripcion}")
    _set_plain(tabla.rows[3].cells[0], f"Datos de Entrenamiento: {datos}")
    _set_plain(tabla.rows[4].cells[0], f"Objetivo / Predicción: {objetivo}")


def _clonar_despues(tabla_modelo, cantidad):
    ancla = tabla_modelo._tbl
    nuevas = []
    for _ in range(cantidad):
        copia = deepcopy(tabla_modelo._tbl)
        ancla.addnext(copia)
        ancla = copia
        nuevas.append(copia)
    return nuevas


def _actualizar_portada(doc):
    for nodo in doc.element.body.iter(qn("w:t")):
        if nodo.text == "PERIODO FEBRERO 2026":
            nodo.text = "PERIODO SEPTIEMBRE 2026"
        elif nodo.text == "Mayo":
            nodo.text = "Septiembre"
        elif nodo.text == " para Requisitos Funcionales (RF)":
            nodo.text = " 3.1 Fichas para Requisitos Funcionales (RF)"
        elif nodo.text == "Fichas para Requisitos Funcionales (RF)":
            nodo.text = "3.1 Fichas para Requisitos Funcionales (RF)"
        elif nodo.text == "Ficha para Requisitos No Funcionales (RNF)":
            nodo.text = "3.2 Fichas para Requisitos No Funcionales (RNF)"
        elif nodo.text == "Ficha para el Componente Inteligente (RCI)":
            nodo.text = "3.3 Fichas para el Componente Inteligente (RCI)"

    for parrafo in doc.element.body.iter(qn("w:p")):
        texto = "".join(parrafo.itertext())
        if "Alvarado" not in texto and "Luis Herice" not in texto:
            continue
        for nodo in parrafo.iter(qn("w:t")):
            if nodo.text in INTEGRANTES_SALIDOS:
                nodo.text = ""


def _quitar_flotante(tabla):
    pr = tabla._tbl.find(qn("w:tblPr"))
    if pr is None:
        return
    ppr = pr.find(qn("w:tblpPr"))
    if ppr is not None:
        pr.remove(ppr)


def _parrafo_espacio():
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    spa = OxmlElement("w:spacing")
    spa.set(qn("w:before"), "120")
    spa.set(qn("w:after"), "120")
    ppr.append(spa)
    p.append(ppr)
    return p


def _no_partir_filas(tabla):
    for tr in tabla._tbl.findall(qn("w:tr")):
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            tr.insert(0, trpr)
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))


def _es_parrafo_vacio(nodo):
    if nodo is None or nodo.tag != qn("w:p"):
        return False
    texto = "".join(nodo.itertext()).strip()
    return texto == ""


def _tablas_en_flujo(doc):
    """Las fichas originales van flotando; al clonar se superponen. Quedan en el flujo."""
    for tabla in doc.tables:
        _quitar_flotante(tabla)
    for tabla in doc.tables[1:]:
        _no_partir_filas(tabla)
    for tabla in doc.tables[1:-1]:
        siguiente = tabla._tbl.getnext()
        while _es_parrafo_vacio(siguiente):
            borrar = siguiente
            siguiente = siguiente.getnext()
            borrar.getparent().remove(borrar)
        tabla._tbl.addnext(_parrafo_espacio())


def generar():
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    copy2(ORIGEN, DESTINO)
    doc = Document(str(DESTINO))

    _actualizar_portada(doc)
    _rellenar_matriz(doc.tables[0])

    for i, ficha in enumerate(RF_FICHAS[:10]):
        _rellenar_rf(doc.tables[1 + i], ficha)

    extra_rf = RF_FICHAS[10:]
    if extra_rf:
        _clonar_despues(doc.tables[10], len(extra_rf))
        for j, ficha in enumerate(extra_rf):
            _rellenar_rf(doc.tables[11 + j], ficha)

    tablas = doc.tables
    desplazamiento = len(extra_rf)
    base_rnf = 11 + desplazamiento
    for i, ficha in enumerate(RNF_FICHAS[:10]):
        _rellenar_rnf(tablas[base_rnf + i], ficha)

    extra_rnf = RNF_FICHAS[10:]
    if extra_rnf:
        _clonar_despues(tablas[base_rnf + 9], len(extra_rnf))
        tablas = doc.tables
        for j, ficha in enumerate(extra_rnf):
            _rellenar_rnf(tablas[base_rnf + 10 + j], ficha)

    desplazamiento_rnf = len(extra_rnf)
    base_rci = base_rnf + 10 + desplazamiento_rnf
    for i, ficha in enumerate(RCI_FICHAS):
        _rellenar_rci(tablas[base_rci + i], ficha)

    val = tablas[base_rci + 3]
    for i, fila in enumerate(VALIDACION_RCI, start=1):
        for j, valor in enumerate(fila):
            _set_plain(val.rows[i].cells[j], valor)

    _tablas_en_flujo(doc)
    doc.save(DESTINO)
    return DESTINO


if __name__ == "__main__":
    print("OK", generar())
