#!/usr/bin/env python3
"""Exporta requisitos markdown a Word y PDF."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown import markdown
from xhtml2pdf import pisa

RAIZ = Path(__file__).resolve().parent
ARCHIVOS = [
    (
        RAIZ / "REQUISITOS_FUNCIONALES.md",
        "Requisitos funcionales — SIGEL / Travel BQTO",
        RAIZ / "REQUISITOS_FUNCIONALES.docx",
        RAIZ / "REQUISITOS_FUNCIONALES.pdf",
    ),
    (
        RAIZ / "REQUISITOS_NO_FUNCIONALES.md",
        "Requisitos no funcionales — SIGEL / Travel BQTO",
        RAIZ / "REQUISITOS_NO_FUNCIONALES.docx",
        RAIZ / "REQUISITOS_NO_FUNCIONALES.pdf",
    ),
]


def _aplicar_fuente(run, tamano=11, negrita=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(tamano)
    run.bold = negrita
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def _agregar_texto_con_negritas(parrafo, texto, tamano=11):
    partes = texto.split("**")
    for i, parte in enumerate(partes):
        if not parte:
            continue
        run = parrafo.add_run(parte)
        _aplicar_fuente(run, tamano=tamano, negrita=(i % 2 == 1))


def md_a_docx(ruta_md: Path, titulo: str, ruta_docx: Path) -> None:
    doc = Document()
    for seccion in doc.sections:
        seccion.top_margin = Cm(2)
        seccion.bottom_margin = Cm(2)
        seccion.left_margin = Cm(2.2)
        seccion.right_margin = Cm(2.2)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    h = doc.add_heading(titulo, level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)

    lineas = ruta_md.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()

        if not linea:
            i += 1
            continue

        if linea.startswith("### "):
            h2 = doc.add_heading(linea[4:].strip(), level=1)
            for run in h2.runs:
                run.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)
            i += 1
            continue

        if linea.startswith("|") and "Campo" in linea:
            filas = []
            i += 1
            if i < len(lineas) and set(lineas[i].replace("|", "").replace("-", "").replace(" ", "")) == set():
                i += 1
            while i < len(lineas) and lineas[i].startswith("|"):
                celdas = [c.strip().replace("**", "") for c in lineas[i].strip().strip("|").split("|")]
                if len(celdas) >= 2:
                    filas.append(celdas[:2])
                i += 1
            tabla = doc.add_table(rows=1, cols=2)
            tabla.style = "Table Grid"
            hdr = tabla.rows[0].cells
            hdr[0].text = "Campo"
            hdr[1].text = "Contenido"
            for celda in hdr:
                for p in celda.paragraphs:
                    for run in p.runs:
                        _aplicar_fuente(run, 10, True)
            for campo, contenido in filas:
                fila = tabla.add_row().cells
                fila[0].text = campo
                fila[1].text = contenido
                for p in fila[0].paragraphs:
                    for run in p.runs:
                        _aplicar_fuente(run, 10, True)
                for p in fila[1].paragraphs:
                    for run in p.runs:
                        _aplicar_fuente(run, 10, False)
            doc.add_paragraph()
            continue

        if linea == "---":
            i += 1
            continue

        if linea.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _agregar_texto_con_negritas(p, linea[2:].strip(), 11)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", linea):
            p = doc.add_paragraph(style="List Number")
            texto = re.sub(r"^\d+\.\s+", "", linea)
            _agregar_texto_con_negritas(p, texto, 11)
            i += 1
            continue

        p = doc.add_paragraph()
        _agregar_texto_con_negritas(p, linea, 11)
        i += 1

    doc.save(ruta_docx)


CSS = """
@page { size: A4; margin: 1.8cm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 11pt; color: #1a1a1a; }
h1 { font-size: 16pt; color: #0d47a1; }
h3 { font-size: 13pt; color: #1565c0; page-break-after: avoid; margin-top: 18pt; }
table { width: 100%; border-collapse: collapse; margin: 8pt 0 12pt 0; }
th, td { border: 1px solid #424242; padding: 6pt 8pt; vertical-align: top; font-size: 10pt; }
th { background: #e3f2fd; }
td:first-child { width: 32%; font-weight: bold; }
ul, ol { margin: 6pt 0 10pt 16pt; }
hr { border: none; border-top: 1px solid #bbb; margin: 16pt 0; }
p { margin: 6pt 0; }
strong { font-weight: bold; }
"""


def md_a_pdf(ruta_md: Path, titulo: str, ruta_pdf: Path) -> None:
    cuerpo = markdown(
        ruta_md.read_text(encoding="utf-8"),
        extensions=["tables", "nl2br"],
    )
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"/><style>{CSS}</style></head>
<body><h1>{titulo}</h1>{cuerpo}</body></html>"""
    with ruta_pdf.open("wb") as destino:
        resultado = pisa.CreatePDF(html, dest=destino, encoding="utf-8")
    if resultado.err:
        raise RuntimeError(f"No se pudo generar PDF: {ruta_pdf}")


def main() -> None:
    for md, titulo, docx, pdf in ARCHIVOS:
        md_a_docx(md, titulo, docx)
        md_a_pdf(md, titulo, pdf)
        print(f"OK {docx.name} / {pdf.name}")


if __name__ == "__main__":
    main()
