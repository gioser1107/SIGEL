#!/usr/bin/env python3
"""SRS en formato UPTAEB Trayecto III (matriz + fichas)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from xhtml2pdf import pisa

RAIZ = Path(__file__).resolve().parent
TITULO_PROYECTO = (
    "DESARROLLO DE UN SISTEMA INTEGRAL PARA LA GESTIÓN LOGÍSTICA TURÍSTICA "
    "CON MÓDULOS DE ANALÍTICA PREDICTIVA Y ARQUITECTURA MVC PARA LA AGENCIA "
    "TRAVEL BQTO, PARROQUIA ANA SOTO, BARQUISIMETO"
)

RF_MATRIZ = [
    ("RF-01", "Selección topológica de asientos", "Funcional", "Alta",
     "Descripción: El sistema debe permitir seleccionar asientos desde un mapa interactivo del vehículo.\n"
     "Entradas: identificador de viaje, identificador de pasajero, número de asiento.\n"
     "Criterio de éxito: el asiento queda bloqueado de forma transaccional; un segundo intento sobre el mismo asiento vigente recibe HTTP 409.",
     "Sobreventa de cupos y disputas físicas el día del viaje."),
    ("RF-02", "Registro y auditoría de pagos", "Funcional", "Alta",
     "Descripción: El sistema debe registrar abonos con método, referencia, comprobante y calcular el saldo deudor.\n"
     "Entradas: monto, referencia, fecha, moneda, método de pago, archivo de comprobante.\n"
     "Criterio de éxito: el pago queda en estado en_validacion (salvo efectivo administrativo) y el saldo se actualiza al aprobar.",
     "Pérdida de referencias bancarias en WhatsApp y errores de cálculo manual."),
    ("RF-03", "Check-in y manifiesto de abordaje", "Funcional", "Alta",
     "Descripción: El sistema debe listar los pasajeros del viaje y registrar asistencia (abordado / no presentado).\n"
     "Entradas: identificador de viaje, identificador de pasajero, estado de asistencia.\n"
     "Criterio de éxito: el manifiesto refleja el estado en tiempo de operación; la anulación devuelve al pasajero a pendiente.",
     "Guías desconectados de la base de datos usando listas de papel."),
    ("RF-04", "Gestión de catálogo de destinos", "Funcional", "Alta",
     "Descripción: El sistema debe permitir crear, editar, anular lógicamente y publicar destinos con galería.\n"
     "Entradas: nombre, descripción, precio base, imágenes.\n"
     "Criterio de éxito: el destino activo se visualiza en el portal público; la anulación se bloquea si hay viajes o reservas dependientes.",
     "El personal transcribe de forma repetitiva la oferta por WhatsApp."),
    ("RF-05", "Solicitud y gestión de cotizaciones", "Funcional", "Alta",
     "Descripción: El sistema debe permitir generar cotizaciones con líneas de desglose, modificarlas y cancelarlas.\n"
     "Entradas: cliente, destino, vigencia, precios, líneas de desglose.\n"
     "Criterio de éxito: los totales se recalculan al modificar líneas; cancelación lógica si no hay reservas activas asociadas.",
     "Latencia para captar prospectos con requerimientos especiales."),
    ("RF-06", "Estructuración de costos operativos", "Funcional", "Alta",
     "Descripción: El sistema debe registrar gastos de logística asociados a un viaje y calcular totales.\n"
     "Entradas: costos de transporte, guías, viáticos y demás partidas del viaje.\n"
     "Criterio de éxito: sumatoria automática del total de gastos y precio de venta asociado al viaje.",
     "Riesgo de pérdida financiera por cálculos empíricos antes de publicar el viaje."),
    ("RF-07", "Planificación de rutas de recogida", "Funcional", "Alta",
     "Descripción: El sistema debe estructurar paradas de embarque y vincular domicilios de clientes al viaje.\n"
     "Entradas: identificador de viaje, paradas, horarios, puntos de recogida del cliente.\n"
     "Criterio de éxito: ruta ordenada consultable en operación; el cliente gestiona sus domicilios en el portal.",
     "Desorganización logística en la recogida de pasajeros."),
    ("RF-08", "Sistema de reseñas y calificaciones", "Funcional", "Media",
     "Descripción: El sistema debe habilitar evaluación posterior al viaje y permitir moderación administrativa.\n"
     "Entradas: identificador de reserva o viaje, puntuación 1–5, comentario.\n"
     "Criterio de éxito: solo reservas elegibles permiten reseña; el administrador puede eliminar reseñas inapropiadas.",
     "Carencia de retroalimentación digital para validar la oferta."),
    ("RF-09", "Validación y aprobación de pagos", "Funcional", "Alta",
     "Descripción: El sistema debe permitir auditar comprobantes, aprobar o rechazar y actualizar el saldo.\n"
     "Entradas: identificador de pago, veredicto, observaciones.\n"
     "Criterio de éxito: aprobación deja estado aprobado y reduce saldo; rechazo no altera saldo ya aprobado; acciones en bitácora.",
     "Falta de conciliación bancaria estructurada y riesgo de fraude."),
    ("RF-10", "Asignación y cronograma de guías", "Funcional", "Media",
     "Descripción: El sistema debe vincular guías a un viaje y limitar la operación de abordaje según rol y permisos.\n"
     "Entradas: identificador de viaje, identificador de usuario con rol guía.\n"
     "Criterio de éxito: el guía autenticado consulta manifiesto y registra abordaje solo con permiso correspondiente.",
     "Desconexión del personal de campo y asignación empírica de guías."),
    ("RF-11", "Registro y gestión de clientes", "Funcional", "Alta",
     "Descripción: El sistema debe permitir crear, consultar, editar y desactivar clientes con documento único.\n"
     "Entradas: tipo y número de documento, nombres, contacto, ubicación.\n"
     "Criterio de éxito: no existe pasajero de reserva sin registro en clientes; documento duplicado se rechaza.",
     "Reservas y acompañantes sin maestro de personas; datos dispersos."),
    ("RF-12", "Gestión de reservas", "Funcional", "Alta",
     "Descripción: El sistema debe crear reservas asociadas a un viaje y un cliente titular, con viajeros en reserva_clientes.\n"
     "Entradas: identificador de viaje, cliente titular, pasajeros, tarifas, puntos de recogida.\n"
     "Criterio de éxito: la reserva debitará cupo; la anulación libera asientos; el cliente consulta solo las suyas en el portal.",
     "Coordinación manual de cupos y listas de pasajeros en canales no formales."),
    ("RF-13", "Portal público de catálogo y agenda", "Funcional", "Alta",
     "Descripción: El sistema debe mostrar destinos y viajes publicados sin autenticación.\n"
     "Entradas: filtros opcionales de destino y fechas.\n"
     "Criterio de éxito: el visitante consulta catálogo y agenda; existen enlaces a registro e inicio de sesión.",
     "Oferta no disponible 24/7 fuera de WhatsApp."),
    ("RF-14", "Portal de autoservicio del cliente", "Funcional", "Alta",
     "Descripción: El sistema debe permitir al cliente autenticado gestionar reservas, pagos, recogida y reseñas.\n"
     "Entradas: credenciales de rol Cliente; datos de la operación solicitada.\n"
     "Criterio de éxito: el cliente no accede a datos de otros clientes; puede reportar pago y completar abono desde el portal.",
     "Dependencia permanente de atención humana para trámites del viajero."),
    ("RF-15", "Reportes estadísticos por rango de fechas", "Funcional", "Alta",
     "Descripción: El sistema debe mostrar indicadores descriptivos filtrados por fechas reales (clientes, reservas, destinos, ingresos aprobados).\n"
     "Entradas: fecha desde, fecha hasta.\n"
     "Criterio de éxito: rangos invertidos o años imposibles retornan HTTP 400; no se reporta por género ni edad (datos no almacenados).",
     "Decisiones de gerencia sin cifras consolidadas del periodo."),
    ("RF-16", "Administración de flota y mapa de asientos", "Funcional", "Alta",
     "Descripción: El sistema debe registrar unidades de transporte y configurar asientos por unidad.\n"
     "Entradas: placa, modelo, capacidad, filas, columnas, tipo de asiento.\n"
     "Criterio de éxito: la unidad queda disponible para asignar a viajes; placa única; baja lógica sin borrar historial.",
     "Mapa de asientos inexistente y cupos no ligados a la flota real."),
    ("RF-17", "Configuración del catálogo financiero", "Funcional", "Alta",
     "Descripción: El sistema debe mantener monedas, tasas (incluida sincronización BCV), métodos de pago, bancos y puntos de venta.\n"
     "Entradas: códigos, valores de tasa, vigencia, estado activo.\n"
     "Criterio de éxito: sin tasa del día, el pago en divisas se rechaza; métodos activos aparecen en el formulario de pago.",
     "Conversión empírica de montos y catálogo de cobro no estandarizado."),
    ("RF-18", "Planificación de viajes", "Funcional", "Alta",
     "Descripción: El sistema debe crear viajes con destino, fechas, unidad, estados y reporte imprimible.\n"
     "Entradas: destino, fechas, unidad, estado operativo.\n"
     "Criterio de éxito: anulación bloqueada si hay reservas activas; el reporte de viaje es consultable e imprimible.",
     "Salidas operativas sin planificación centralizada en un solo registro de viaje."),
    ("RF-19", "Consulta de bitácora de auditoría", "Funcional", "Alta",
     "Descripción: El sistema debe permitir consultar eventos de auditoría con filtros.\n"
     "Entradas: módulo, acción, usuario, rango de fechas, texto de búsqueda.\n"
     "Criterio de éxito: sin permiso leer_bitacora la consulta retorna HTTP 403; cada evento muestra actor, módulo, acción y registro.",
     "No hay evidencia de quién modificó pagos, reservas o clientes."),
]

RNF_MATRIZ = [
    ("RNF-01", "Autenticación e inicio de sesión", "No Funcional", "Alta",
     "Categoría: Seguridad / acceso.\n"
     "Descripción: El sistema exigirá credenciales para funciones restringidas.\n"
     "Especificación: Token JWT (Bearer, HS256) con claims sub, exp, iat; expiración configurable ≤ 24 h; usuarios con eliminado_en no autentican.\n"
     "Criterio de éxito: 100 % de rutas administrativas, de guía y de portal cliente rechazan petición sin token válido (HTTP 401).",
     "Operación sin control centralizado de acceso; riesgo de exposición de pagos y pasajeros."),
    ("RNF-02", "Gestión de cuentas de usuario", "No Funcional", "Alta",
     "Categoría: Operación / administración.\n"
     "Descripción: Altas, edición y desactivación de cuentas internas y asociadas a clientes.\n"
     "Especificación: Eliminación lógica; cuentas inactivas no autentican; cambios auditables en bitácora; propagación en la siguiente solicitud autenticada ≤ 60 s.\n"
     "Criterio de éxito: usuario desactivado recibe HTTP 401 al iniciar sesión; no aparece en listados operativos.",
     "Personal segmentado gestionado de forma manual y no estructurada."),
    ("RNF-03", "Navegabilidad y estructura de la interfaz", "No Funcional", "Media",
     "Categoría: Usabilidad / navegación.\n"
     "Descripción: Jerarquía de pantallas predecible por rol.\n"
     "Especificación: Máximo 4 clics desde el inicio del rol hasta catálogo, reserva, mapa de asientos y pagos pendientes; menú persistente.\n"
     "Criterio de éxito: rutas críticas verificadas en plan de pruebas; 0 callejones sin salida en reserva/pago.",
     "Flujo dependiente de mensajes sueltos en lugar de un portal guiado."),
    ("RNF-04", "Portabilidad y compatibilidad cliente", "No Funcional", "Alta",
     "Categoría: Portabilidad.\n"
     "Descripción: Acceso desde escritorio y móviles.\n"
     "Especificación: Interfaz adaptable; Chrome y Edge (dos últimas versiones estables); ancho mínimo 360 px sin pérdida de reserva ni check-in.\n"
     "Criterio de éxito: lista de verificación en dispositivos definidos en el plan de pruebas.",
     "Guías en campo desconectados de listas impresas."),
    ("RNF-05", "Usabilidad en tareas críticas", "No Funcional", "Media",
     "Categoría: Usabilidad.\n"
     "Descripción: Reducción de errores en reserva, asiento y pago.\n"
     "Especificación: Prueba con 5 usuarios; tasa de error ≤ 10 %; reserva de prueba ≤ 10 min sin ayuda; mensajes en español.\n"
     "Criterio de éxito: SUS ≥ 70 y registro cuantitativo en informe de usabilidad.",
     "Sobreventa y errores de cobro por proceso manual."),
    ("RNF-06", "Mantenibilidad del software", "No Funcional", "Media",
     "Categoría: Mantenibilidad.\n"
     "Descripción: Arquitectura que localiza cambios por módulo.\n"
     "Especificación: MVC (controladores, modelos, utilidades); frontend React y API FastAPI; OpenAPI en /docs; objetivo de cobertura ≥ 60 % en reservas, pagos y autenticación.\n"
     "Criterio de éxito: un endpoint CRUD nuevo en ≤ 4 h; informe de cobertura cuando se ejecute la batería de pruebas.",
     "Necesidad de incorporar analítica predictiva sin reescribir el núcleo."),
    ("RNF-07", "Eficiencia y uso de recursos", "No Funcional", "Media",
     "Categoría: Utilización eficiente de los recursos.\n"
     "Descripción: Limitar CPU, memoria y ancho de banda en operación nominal.\n"
     "Especificación: Bundle principal ≤ 500 KB comprimido; paginación en listados > 50 registros; imágenes WebP; RAM del proceso API ≤ 512 MB con 25 usuarios.\n"
     "Criterio de éxito: medición en build y revisión de consultas paginadas.",
     "Infraestructura limitada y alto volumen transaccional relativo."),
    ("RNF-08", "Fiabilidad e integridad transaccional", "No Funcional", "Alta",
     "Categoría: Fiabilidad.\n"
     "Descripción: Operaciones críticas atómicas sobre cupos, asientos y pagos.\n"
     "Especificación: Transacciones ACID MySQL InnoDB; conflicto de integridad HTTP 409; fallo de BD HTTP 503 sin traza al usuario.\n"
     "Criterio de éxito: 0 sobreventas en prueba de reservas concurrentes sobre el mismo asiento.",
     "Conflictos de concurrencia y sobreventa."),
    ("RNF-09", "Seguridad de datos y comunicaciones", "No Funcional", "Alta",
     "Categoría: Seguridad.\n"
     "Descripción: Confidencialidad en tránsito y reposo.\n"
     "Especificación: TLS 1.2+ en producción; hash unidireccional de contraseñas (prototipo SHA-256; objetivo bcrypt/Argon2; no MD5); secretos fuera del repositorio; uploads con lista blanca MIME y máximo 10 MB.\n"
     "Criterio de éxito: 0 contraseñas en texto plano; recursos protegidos no accesibles de forma anónima.",
     "Comprobantes y referencias por canales no seguros."),
    ("RNF-10", "Validación de datos de entrada", "No Funcional", "Alta",
     "Categoría: Calidad de datos.\n"
     "Descripción: Rechazo de entradas inválidas en cliente y servidor.\n"
     "Especificación: Esquemas Pydantic y reglas de dominio; HTTP 422 o 400; montos ≥ 0; fechas coherentes; nombres sin dígitos.\n"
     "Criterio de éxito: 100 % de endpoints de escritura con payload inválido no persisten el registro.",
     "Errores en pagos fraccionados y omisiones de datos."),
    ("RNF-11", "Permisos", "No Funcional", "Alta",
     "Categoría: Seguridad / autorización.\n"
     "Descripción: Cada acción administrativa se autoriza por permiso granular.\n"
     "Especificación: Permisos crear_*, leer_*, editar_*, borrar_* validados en servidor; HTTP 403 si falta el permiso.\n"
     "Criterio de éxito: un cliente no ejecuta endpoints de conciliación, usuarios ni bitácora.",
     "Acciones de gerencia ejecutables por cualquier cuenta autenticada."),
    ("RNF-12", "Roles", "No Funcional", "Alta",
     "Categoría: Seguridad / autorización.\n"
     "Descripción: El acceso se agrupa en roles (Administrador, Guía, Cliente y equivalentes).\n"
     "Especificación: Tablas roles y roles_permisos; el token incluye rol_id verificado contra la base de datos.\n"
     "Criterio de éxito: cambio de permisos del rol se refleja en la siguiente autenticación.",
     "No existe separación de responsabilidades entre oficina, campo y cliente."),
    ("RNF-13", "Módulos", "No Funcional", "Media",
     "Categoría: Arquitectura.\n"
     "Descripción: El software se organiza en módulos de dominio acoplados por API REST.\n"
     "Especificación: Prefijo /api; routers por dominio (auth, destinos, reservas, pagos, abordaje, reportes, entre otros); OpenAPI.\n"
     "Criterio de éxito: cada módulo aparece documentado en /docs; desactivar un router no impide arrancar el resto.",
     "Riesgo de monolito no separable para mantenimiento académico y evolución."),
    ("RNF-14", "Bitácora", "No Funcional", "Alta",
     "Categoría: Operabilidad / cumplimiento.\n"
     "Descripción: Las acciones críticas quedan registradas.\n"
     "Especificación: Registro con usuario, módulo, acción, tabla, identificador, resumen, detalle, IP y fecha; consulta paginada.\n"
     "Criterio de éxito: login más 3 escrituras generan ≥ 4 registros coherentes.",
     "Imposibilidad de auditar quién alteró un pago o una reserva."),
    ("RNF-15", "Tiempo de respuesta", "No Funcional", "Media",
     "Categoría: Rendimiento.\n"
     "Descripción: Operaciones frecuentes responden dentro de umbral medible.\n"
     "Especificación: Percentil 95 de lecturas ≤ 2 s; escrituras ≤ 3 s; salud de /api ≤ 500 ms.\n"
     "Criterio de éxito: prueba de 20 usuarios virtuales durante 3 minutos con promedio < 2 s y 0 % HTTP 500.",
     "Espera excesiva en portal y panel durante picos de consulta."),
    ("RNF-16", "Escalabilidad", "No Funcional", "Media",
     "Categoría: Escalabilidad.\n"
     "Descripción: Soportar usuarios concurrentes sin degradar integridad.\n"
     "Especificación: JWT sin estado de servidor; paginación; índices; compatible con ≥ 2 procesos de API.\n"
     "Criterio de éxito: 50 usuarios virtuales / 10 min con tasa de error 5xx < 1 % y P95 ≤ 4 s.",
     "Picos de reservas y abordaje simultáneo en temporada alta."),
    ("RNF-17", "Disponibilidad", "No Funcional", "Media",
     "Categoría: Disponibilidad.\n"
     "Descripción: Portal y API accesibles en horario operativo extendido.\n"
     "Especificación: Verificación de salud en /api; respaldo diario de MySQL con retención de 7 días; objetivo de disponibilidad mensual ≥ 99 % en despliegue de producción.\n"
     "Criterio de éxito: restauración de respaldo < 60 min; mensaje controlado ante caída de BD (HTTP 503).",
     "La operación turística requiere acceso fuera del horario de oficina."),
    ("RNF-18", "Rendimiento", "No Funcional", "Media",
     "Categoría: Rendimiento / carga.\n"
     "Descripción: El sistema mantiene tasa de error y latencia acotadas bajo carga nominal.\n"
     "Especificación: 25 usuarios concurrentes, 70 % lectura y 30 % escritura; medición con Apache JMeter o Locust.\n"
     "Criterio de éxito: P95 dentro de RNF-15; sin errores HTTP 500 atribuibles a agotamiento de recursos en la prueba documentada.",
     "Alto volumen relativo de consultas de catálogo y reservas en periodos de venta."),
]

RF_FICHAS = [
    ("RF-01", "Selección topológica de asientos",
     "El sistema debe renderizar el mapa de asientos del viaje y permitir reservar un puesto específico sin sobreventa.",
     "Cliente, Administrador, Atención al cliente.",
     "Identificador de viaje, identificador de pasajero, número de asiento.",
     "Asiento vigente ocupado; respuesta HTTP 409 si ya está tomado; comprobante de asignación en la reserva.",
     "El viaje debe existir y el asiento debe estar disponible; el actor debe estar autenticado con permiso o ser el cliente titular."),
    ("RF-02", "Registro y auditoría de pagos",
     "El sistema debe capturar abonos, equivalencia según tasa del día y dejar el pago en validación.",
     "Cliente, Administrador.",
     "Monto, referencia, moneda, método, banco o punto de venta, comprobante.",
     "Pago en estado en_validacion; saldo pendiente visible; alerta en bandeja administrativa.",
     "Debe existir una reserva con saldo pendiente; en divisas debe existir tasa del día."),
    ("RF-03", "Check-in y manifiesto de abordaje",
     "El sistema debe proveer el manifiesto digital y registrar abordado o no presentado, individual o por lote.",
     "Guía, Administrador.",
     "Identificador de viaje, identificador de pasajero, estado de asistencia.",
     "Manifiesto actualizado; totales de abordados y pendientes; evento en bitácora.",
     "El usuario debe tener permiso de abordaje; el viaje debe tener pasajeros en reserva."),
    ("RF-04", "Gestión de catálogo de destinos",
     "El sistema debe permitir administrar destinos e imágenes y publicarlos en el portal.",
     "Administrador. El visitante solo consulta.",
     "Nombre, descripción, precio, archivos de imagen.",
     "Destino persistido; galería WebP; visibilidad en catálogo público si está activo.",
     "El usuario debe haber iniciado sesión con permiso de destinos."),
    ("RF-05", "Solicitud y gestión de cotizaciones",
     "El sistema debe permitir crear cotizaciones con líneas, recalcular totales y cancelar lógicamente.",
     "Administrador, Atención al cliente.",
     "Cliente, destino, vigencia, líneas de precio.",
     "Cotización numerada con desglose y estado (solicitada, pendiente, aceptada, rechazada, cancelada).",
     "El actor debe tener permiso de cotizaciones; el cliente y el destino deben existir."),
    ("RF-06", "Estructuración de costos operativos",
     "El sistema debe asentar gastos del viaje y calcular total y precio de venta.",
     "Administrador.",
     "Partidas de costo asociadas al identificador de viaje.",
     "Totales almacenados y visibles en la ficha y el reporte del viaje.",
     "Debe existir el viaje en planificación o estado editable."),
    ("RF-07", "Planificación de rutas de recogida",
     "El sistema debe ordenar paradas del viaje y asociar puntos de recogida de cada cliente.",
     "Administrador, Cliente (domicilios propios).",
     "Paradas, horarios, identificadores de puntos de recogida.",
     "Ruta ordenada; domicilios del cliente persistidos en clientes_puntos_recogida.",
     "El viaje debe existir; el cliente debe estar registrado para vincular domicilio."),
    ("RF-08", "Sistema de reseñas y calificaciones",
     "El sistema debe desbloquear la reseña para reservas elegibles y permitir moderación.",
     "Cliente, Administrador.",
     "Puntuación 1–5, comentario opcional, identificador de viaje o reserva.",
     "Reseña almacenada; listado público; eliminación lógica o física según regla administrativa.",
     "La reserva debe ser elegible (viaje o pago en condición definida por el módulo de reseñas)."),
    ("RF-09", "Validación y aprobación de pagos",
     "El sistema debe permitir aprobar o rechazar reportes de pago y actualizar el saldo.",
     "Administrador.",
     "Identificador de pago, acción aprobar o rechazar, observación.",
     "Estado aprobado o rechazado; saldo de reserva actualizado; bitácora.",
     "El pago debe existir en estado en_validacion (RF-02)."),
    ("RF-10", "Asignación y cronograma de guías",
     "El sistema debe asignar guías al viaje y habilitar operación de manifiesto según rol.",
     "Administrador, Guía.",
     "Identificador de viaje, identificador de usuario guía.",
     "Vinculación persistida; el guía consulta viajes y manifiesto autorizados.",
     "El usuario debe tener rol Guía activo; el viaje debe estar estructurado (RF-18 / RF-04)."),
    ("RF-11", "Registro y gestión de clientes",
     "El sistema debe mantener el maestro de clientes con documento único y desactivación lógica.",
     "Administrador.",
     "Tipo y número de documento, nombres, teléfono, ubicación.",
     "Cliente persistido; búsqueda por documento; no se destruye historial al desactivar.",
     "El actor debe tener permiso de clientes."),
    ("RF-12", "Gestión de reservas",
     "El sistema debe crear la reserva con titular y viajeros, cada viajero siendo un cliente.",
     "Administrador, Cliente.",
     "Viaje, cliente titular, lista de pasajeros, tarifas, recogida.",
     "Reserva con estados pendiente, confirmada, abonada o cancelada; cupo actualizado.",
     "El viaje debe tener cupo; los pasajeros deben existir como clientes (RF-11)."),
    ("RF-13", "Portal público de catálogo y agenda",
     "El sistema debe publicar destinos y viajes consultables sin autenticación.",
     "Visitante.",
     "Filtros opcionales de destino y fechas.",
     "Listado y detalle de destinos; agenda de viajes publicados.",
     "Ninguna. Es un flujo público."),
    ("RF-14", "Portal de autoservicio del cliente",
     "El sistema debe ofrecer al cliente autenticado sus reservas, pagos, recogida y reseñas.",
     "Cliente.",
     "Token de sesión de rol Cliente y datos de la operación.",
     "Vistas restringidas a sus registros; pago en_validacion; domicilios propios.",
     "El usuario debe haberse autenticado (RNF-01) con rol Cliente."),
    ("RF-15", "Reportes estadísticos por rango de fechas",
     "El sistema debe calcular indicadores descriptivos del periodo indicado.",
     "Administrador.",
     "Fecha desde, fecha hasta.",
     "Totales de clientes, reservas, pasajeros, destinos e ingresos de pagos aprobados; opción de impresión.",
     "El actor debe tener al menos uno de los permisos de lectura de reservas, pagos o clientes."),
    ("RF-16", "Administración de flota y mapa de asientos",
     "El sistema debe registrar unidades y su mapa de asientos.",
     "Administrador.",
     "Placa, modelo, capacidad, geometría de asientos.",
     "Unidad y asientos persistidos; disponibles para RF-18 y RF-01.",
     "Sesión con permiso de flota o asientos."),
    ("RF-17", "Configuración del catálogo financiero",
     "El sistema debe mantener monedas, tasas, métodos, bancos y puntos de venta.",
     "Administrador.",
     "Valores de catálogo y tasa del día.",
     "Catálogos activos utilizables en RF-02; error si falta tasa requerida.",
     "Sesión con permiso de catálogo de pagos."),
    ("RF-18", "Planificación de viajes",
     "El sistema debe crear y administrar viajes programados y su reporte.",
     "Administrador.",
     "Destino, fechas, unidad, estado, costos (RF-06), guías (RF-10).",
     "Viaje persistido; reporte imprimible; anulación bloqueada con reservas activas.",
     "Deben existir destino (RF-04) y unidad (RF-16) cuando la regla de negocio lo exija."),
    ("RF-19", "Consulta de bitácora de auditoría",
     "El sistema debe listar eventos de auditoría filtrables.",
     "Administrador.",
     "Filtros de módulo, acción, usuario y fechas.",
     "Listado paginado con actor, módulo, acción, registro e IP.",
     "Permiso leer_bitacora (RNF-14)."),
]

RNF_FICHAS = [
    ("RNF-01", "Autenticación e inicio de sesión",
     "El sistema deberá exigir autenticación para funciones restringidas e invalidar tokens expirados.",
     "Porcentaje de rutas protegidas que retornan HTTP 401 sin token; revisión de expiración JWT.",
     "0 accesos no autenticados a usuarios, reservas administrativas, pagos y bitácora.",
     "JWT Bearer HS256; verificación de usuario y rol activos; HTTPS en producción.",
     "Alta"),
    ("RNF-02", "Gestión de cuentas de usuario",
     "El sistema deberá permitir crear, editar y desactivar cuentas sin borrar el historial legal mínimo.",
     "Cuentas inactivas no autentican; tiempo de efecto del cambio ≤ 60 s en la API.",
     "Listados operativos sin cuentas desactivadas; bitácora del cambio de estado.",
     "Eliminación lógica (eliminado_en); estados de cuenta; auditoría.",
     "Alta"),
    ("RNF-03", "Navegabilidad y estructura de la interfaz",
     "El sistema deberá ofrecer navegación estable por rol hasta las tareas turísticas críticas.",
     "Profundidad ≤ 4 clics en catálogo, reserva, asientos y pagos pendientes.",
     "Rutas reproducibles en el plan de pruebas, sin callejón sin salida.",
     "Menú por rol; mapa de sitio portal público, portal cliente y back-office.",
     "Media"),
    ("RNF-04", "Portabilidad y compatibilidad cliente",
     "El sistema deberá operar en escritorio y móvil para consulta y check-in.",
     "Chrome y Edge (dos últimas versiones); viewport 360 px.",
     "Reserva y manifiesto utilizables en móvil sin pérdida funcional.",
     "CSS adaptable; pruebas de resolución documentadas.",
     "Alta"),
    ("RNF-05", "Usabilidad en tareas críticas",
     "El sistema deberá permitir completar reserva, asiento y pago con error acotado.",
     "5 usuarios; error ≤ 10 %; tarea ≤ 10 min; SUS ≥ 70.",
     "Informe cuantitativo de usabilidad.",
     "Mensajes en español; validación en línea coherente con RNF-10.",
     "Media"),
    ("RNF-06", "Mantenibilidad del software",
     "El sistema deberá localizar cambios por capa MVC y por módulo de API.",
     "Cobertura objetivo ≥ 60 % en autenticación, reservas y pagos; revisión de carpetas controladores/modelos/utilidades.",
     "OpenAPI disponible; incorporación de módulos sin reescritura del núcleo.",
     "FastAPI + React; convenciones de código.",
     "Media"),
    ("RNF-07", "Eficiencia y uso de recursos",
     "El sistema deberá limitar consumo de red y servidor en operación nominal.",
     "Bundle ≤ 500 KB comprimido; paginación > 50 filas; RAM API ≤ 512 MB / 25 usuarios.",
     "Métricas de build y consultas dentro de umbral.",
     "Paginación, índices y compresión de imágenes WebP.",
     "Media"),
    ("RNF-08", "Fiabilidad e integridad transaccional",
     "El sistema deberá conservar consistencia ante concurrencia y fallo de base de datos.",
     "0 sobreventas en prueba concurrente del mismo asiento; HTTP 503 genérico si cae MySQL.",
     "Inventario de asientos coherente; sin traza SQL al usuario.",
     "ACID InnoDB; HTTP 409 en conflicto de integridad.",
     "Alta"),
    ("RNF-09", "Seguridad de datos y comunicaciones",
     "El sistema deberá proteger credenciales y evidencias de pago.",
     "TLS 1.2+ en producción; 0 contraseñas en texto plano; 401/403 en API sensible sin token.",
     "Hash unidireccional; comprobantes con tipo MIME restringido.",
     "HTTPS; SHA-256 en prototipo con objetivo bcrypt/Argon2; .env fuera de git.",
     "Alta"),
    ("RNF-10", "Validación de datos de entrada",
     "El sistema deberá rechazar entradas inválidas antes de persistir.",
     "100 % de POST/PUT con payload inválido → HTTP 422 o 400 sin inserción.",
     "Mensaje controlado sin datos internos de infraestructura.",
     "Pydantic y validador de dominio en servidor; validación complementaria en cliente.",
     "Alta"),
    ("RNF-11", "Permisos",
     "El sistema deberá autorizar cada operación administrativa mediante permiso explícito.",
     "100 % de intentos de Cliente sobre endpoints de conciliación o usuarios → HTTP 403.",
     "Matriz rol-permiso aplicable en servidor, no solo en menú.",
     "Dependencias requiere_permiso y requiere_alguno_de_permisos.",
     "Alta"),
    ("RNF-12", "Roles",
     "El sistema deberá agrupar permisos en roles persistidos.",
     "Rol Cliente no lista ni modifica catálogos administrativos; Guía limitado a abordaje según permiso.",
     "roles y roles_permisos consistentes con el token.",
     "RBAC; rol_id verificado en cada solicitud.",
     "Alta"),
    ("RNF-13", "Módulos",
     "El sistema deberá exponer capacidades por módulos REST independientes.",
     "Routers documentados en OpenAPI /docs.",
     "Fallo de un módulo de negocio no impide arrancar autenticación y catálogo público.",
     "Prefijo /api; un controlador por dominio.",
     "Media"),
    ("RNF-14", "Bitácora",
     "El sistema deberá registrar actor, módulo, acción y registro afectado en operaciones críticas.",
     "Tras login y 3 escrituras, ≥ 4 filas coherentes; consulta sin permiso → 403.",
     "Trazabilidad de pagos, reservas, clientes y abordaje.",
     "Tabla bitacora con IP, detalle y fecha.",
     "Alta"),
    ("RNF-15", "Tiempo de respuesta",
     "El sistema deberá responder lecturas frecuentes en ≤ 2 s (percentil 95).",
     "JMeter o Locust: 20 usuarios, 3 minutos, promedio < 2 s, 0 % HTTP 500.",
     "Salud /api ≤ 500 ms en vacío de carga.",
     "Medición sobre catálogo público y API local.",
     "Media"),
    ("RNF-16", "Escalabilidad",
     "El sistema deberá admitir incremento de usuarios concurrentes sin perder integridad.",
     "50 usuarios virtuales / 10 min; 5xx < 1 %; P95 ≤ 4 s.",
     "Sesión basada en JWT; listados paginados.",
     "Más de un proceso de API; índices en claves de búsqueda.",
     "Media"),
    ("RNF-17", "Disponibilidad",
     "El sistema deberá permanecer operable o fallar de forma controlada.",
     "Objetivo ≥ 99 % mensual en producción; restauración < 60 min; 503 controlado si cae MySQL.",
     "Respaldo diario retención 7 días; endpoint de salud /api.",
     "Monitoreo académico en local mediante prueba de interrupción de BD.",
     "Media"),
    ("RNF-18", "Rendimiento",
     "El sistema deberá sostener carga nominal sin errores 500 por agotamiento.",
     "25 usuarios; 70 % lectura / 30 % escritura; P95 alineado a RNF-15.",
     "Informe de Locust o JMeter con RPS y latencia.",
     "Prueba sobre /api y /api/catalogo/destinos.",
     "Media"),
]


def _arial(run, size=12, bold=False, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def _p(doc, texto, size=12, bold=False, center=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(texto)
    _arial(run, size, bold)
    return p


def _sombrear(celda, hex_color):
    tc = celda._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _celda(celda, texto, bold=False, size=10):
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(texto)
    _arial(run, size, bold)


def _tabla_matriz(doc, filas):
    tabla = doc.add_table(rows=1, cols=6)
    tabla.style = "Table Grid"
    headers = [
        "ID",
        "Nombre del Requisito",
        "Tipo",
        "Prioridad",
        "Especificación Técnica y Criterios de Aceptación",
        "Necesidad o Problema Asociado",
    ]
    for i, h in enumerate(headers):
        _celda(tabla.rows[0].cells[i], h, bold=True, size=10)
        _sombrear(tabla.rows[0].cells[i], "D6EAF8")
    for fila in filas:
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            _celda(celdas[i], valor, bold=(i == 0), size=10)
    for fila in tabla.rows:
        for celda in fila.cells:
            for p in celda.paragraphs:
                p.paragraph_format.space_after = Pt(2)


def _ficha_rf(doc, ficha):
    ident, nombre, desc, actores, entradas, salidas, pre = ficha
    _p(doc, f"ID: {ident}", bold=True, space_after=2)
    _p(doc, f"Nombre: {nombre}", space_after=2)
    _p(doc, f"Descripción: {desc}", space_after=2)
    _p(doc, f"Actores: {actores}", space_after=2)
    _p(doc, f"Entradas: {entradas}", space_after=2)
    _p(doc, f"Salidas: {salidas}", space_after=2)
    _p(doc, f"Precondición: {pre}", space_after=12)


def _ficha_rnf(doc, ficha):
    ident, nombre, desc, metrica, resultados, espec, prioridad = ficha
    _p(doc, f"ID: {ident}", bold=True, space_after=2)
    _p(doc, f"Nombre: {nombre}", space_after=2)
    _p(doc, f"Descripción: {desc}", space_after=2)
    _p(doc, f"Métrica y evaluación: {metrica}", space_after=2)
    _p(doc, f"Resultados: {resultados}", space_after=2)
    _p(doc, f"Especificación: {espec}", space_after=2)
    _p(doc, f"Prioridad: {prioridad}", space_after=12)


def _portada(doc, subtitulo):
    _p(doc, "PROGRAMA NACIONAL DE FORMACIÓN EN INFORMÁTICA", 14, True, True, 8)
    _p(doc, TITULO_PROYECTO, 12, True, True, 8)
    _p(doc, "Especificación de Requisitos de Software (SRS)", 12, True, True, 8)
    _p(doc, subtitulo, 12, True, True, 12)
    _p(doc, "Integrantes: María Alvarado, Luis Herice, Sergio Jiménez, Gabriel Jiménez", 12, False, True, 4)
    _p(doc, "Tutor: Edecio Freitez", 12, False, True, 4)
    _p(doc, "Barquisimeto, septiembre de 2026", 12, False, True, 16)


def _setup(doc):
    for seccion in doc.sections:
        seccion.top_margin = Cm(2)
        seccion.bottom_margin = Cm(2)
        seccion.left_margin = Cm(1.8)
        seccion.right_margin = Cm(1.8)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(12)


def generar_rf():
    doc = Document()
    _setup(doc)
    _portada(doc, "Requisitos funcionales (RF)")
    _p(doc, "2. Matriz Única de Especificación de Requisitos (SRS)", 12, True, False, 8)
    _tabla_matriz(doc, RF_MATRIZ)
    doc.add_page_break()
    _p(doc, "3. Fichas Técnicas de Detalle — Requisitos Funcionales (RF)", 12, True, False, 10)
    _p(doc, "Campos obligatorios: ID, Nombre, Descripción, Actores, Entradas, Salidas, Precondición.", 12, False, False, 10)
    for ficha in RF_FICHAS:
        _ficha_rf(doc, ficha)
    ruta = RAIZ / "REQUISITOS_FUNCIONALES.docx"
    doc.save(ruta)
    return ruta


def generar_rnf():
    doc = Document()
    _setup(doc)
    _portada(doc, "Requisitos no funcionales (RNF)")
    _p(doc, "Los RNF cubren las categorías obligatorias: inicio de sesión, gestión de usuarios, navegabilidad, portabilidad, usabilidad, mantenibilidad, eficiencia, fiabilidad, seguridad, validación de entradas, permisos, roles, módulos, bitácora, tiempo de respuesta, escalabilidad, disponibilidad, rendimiento y uso eficiente de recursos.", 12, False, False, 10)
    _p(doc, "2. Matriz Única de Especificación de Requisitos (SRS)", 12, True, False, 8)
    _tabla_matriz(doc, RNF_MATRIZ)
    doc.add_page_break()
    _p(doc, "3. Fichas Técnicas de Detalle — Requisitos No Funcionales (RNF)", 12, True, False, 10)
    _p(doc, "Campos obligatorios: ID, Nombre, Descripción, Métrica y evaluación, Resultados, Especificación, Prioridad.", 12, False, False, 10)
    for ficha in RNF_FICHAS:
        _ficha_rnf(doc, ficha)
    ruta = RAIZ / "REQUISITOS_NO_FUNCIONALES.docx"
    doc.save(ruta)
    return ruta


def docx_a_html_simple(titulo, matriz, fichas, tipo):
    filas = ""
    for f in matriz:
        celdas = "".join(f"<td>{str(c).replace(chr(10), '<br/>')}</td>" for c in f)
        filas += f"<tr>{celdas}</tr>"
    bloque_fichas = ""
    if tipo == "rf":
        for ident, nombre, desc, actores, entradas, salidas, pre in fichas:
            bloque_fichas += f"""
            <h3>ID: {ident}</h3>
            <p><b>Nombre:</b> {nombre}</p>
            <p><b>Descripción:</b> {desc}</p>
            <p><b>Actores:</b> {actores}</p>
            <p><b>Entradas:</b> {entradas}</p>
            <p><b>Salidas:</b> {salidas}</p>
            <p><b>Precondición:</b> {pre}</p>
            """
    else:
        for ident, nombre, desc, metrica, resultados, espec, prioridad in fichas:
            bloque_fichas += f"""
            <h3>ID: {ident}</h3>
            <p><b>Nombre:</b> {nombre}</p>
            <p><b>Descripción:</b> {desc}</p>
            <p><b>Métrica y evaluación:</b> {metrica}</p>
            <p><b>Resultados:</b> {resultados}</p>
            <p><b>Especificación:</b> {espec}</p>
            <p><b>Prioridad:</b> {prioridad}</p>
            """
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8"/>
    <style>
    @page {{ size: A4 landscape; margin: 1.2cm; }}
    body {{ font-family: Helvetica, Arial, sans-serif; font-size: 10pt; }}
    h1,h2 {{ color: #1a365d; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 8pt; }}
    th, td {{ border: 1px solid #333; padding: 4px; vertical-align: top; }}
    th {{ background: #d6eaf8; }}
    h3 {{ margin-top: 14pt; }}
    </style></head><body>
    <h1>PROGRAMA NACIONAL DE FORMACIÓN EN INFORMÁTICA</h1>
    <p><b>{TITULO_PROYECTO}</b></p>
    <h2>{titulo}</h2>
    <p>Integrantes: María Alvarado, Luis Herice, Sergio Jiménez, Gabriel Jiménez. Tutor: Edecio Freitez. Barquisimeto, septiembre de 2026.</p>
    <h2>2. Matriz Única de Especificación de Requisitos (SRS)</h2>
    <table><tr>
    <th>ID</th><th>Nombre del Requisito</th><th>Tipo</th><th>Prioridad</th>
    <th>Especificación Técnica y Criterios de Aceptación</th>
    <th>Necesidad o Problema Asociado</th></tr>
    {filas}</table>
    <h2>3. Fichas Técnicas de Detalle</h2>
    {bloque_fichas}
    </body></html>"""


def generar_pdfs():
    pares = [
        (
            RAIZ / "REQUISITOS_FUNCIONALES.pdf",
            "Requisitos funcionales (RF) — SRS Trayecto III",
            RF_MATRIZ,
            RF_FICHAS,
            "rf",
        ),
        (
            RAIZ / "REQUISITOS_NO_FUNCIONALES.pdf",
            "Requisitos no funcionales (RNF) — SRS Trayecto III",
            RNF_MATRIZ,
            RNF_FICHAS,
            "rnf",
        ),
    ]
    for ruta, titulo, matriz, fichas, tipo in pares:
        html = docx_a_html_simple(titulo, matriz, fichas, tipo)
        with ruta.open("wb") as dest:
            r = pisa.CreatePDF(html, dest=dest, encoding="utf-8")
        if r.err:
            raise RuntimeError(ruta)


if __name__ == "__main__":
    generar_rf()
    generar_rnf()
    generar_pdfs()
    print("OK")
