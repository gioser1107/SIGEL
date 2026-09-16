#!/usr/bin/env python3
"""Carga la planilla real de Travel BQTO (agosto 2026) en SIGEL.

Uso, desde backend/:

  python importar_planilla_agosto.py
  python importar_planilla_agosto.py --forzar
  python importar_planilla_agosto.py --excel ../instalacion/datos_agencia/AGOSTO_TRAVELBQTO.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
import pymysql

directorio_backend = Path(__file__).resolve().parent
load_dotenv(directorio_backend / ".env")

BD_SEGURIDAD = (os.getenv("DB_NOMBRE_SEGURIDAD") or "travel_bqto_seguridad").strip()
if not re.fullmatch(r"[A-Za-z0-9_]+", BD_SEGURIDAD):
    raise SystemExit(f"DB_NOMBRE_SEGURIDAD no válido: {BD_SEGURIDAD}")

EXCEL_DEFECTO = (
    directorio_backend.parent / "instalacion" / "datos_agencia" / "AGOSTO_TRAVELBQTO.xlsx"
)
HASH_CLAVE = hashlib.sha256(b"TravelBqto2026").hexdigest()
MARCA = "Planilla operativa Travel BQTO, agosto 2026"
CI_INTERNO_INICIO = 89000001
ESTADO_LARA = 13
CIUDAD_BQTO = 19
ADMIN_ID = 1
ROL_GUIA = 2

BANCOS_NUEVOS = [
    ("0114", "Bancaribe"),
    ("0174", "Banplus"),
    ("0172", "Bancamiga"),
    ("0191", "Banco Nacional de Crédito"),
    ("0163", "Banco del Tesoro"),
]

ALIAS_BANCO = {
    "VZLA": "0102",
    "VENEZUELA": "0102",
    "BANESCO": "0134",
    "BNAESCO": "0134",
    "MERCANTIL": "0105",
    "PROVINCIAL": "0108",
    "BANCARIBE": "0114",
    "BANPLUS": "0174",
    "BANCAMIGA": "0172",
    "BNC": "0191",
    "TESORO": "0163",
    "BDT": "0163",
}

UNIDADES = [
    ("TBQ-UNI-01", "Unidad I", 31),
    ("TBQ-UNI-02", "Unidad II", 31),
    ("TBQ-UNI-03", "Unidad III", 31),
]

VIAJES = [
    {
        "hoja": "CAYO MUERTO (01-08)",
        "cobranza": "COBRANZA CAYO MUERTO (01-08)",
        "destino": "Cayo Muerto",
        "descripcion": "Full day a Cayo Muerto, Parque Nacional Morrocoy. Salida desde Barquisimeto.",
        "precio": Decimal("23.00"),
        "dificultad": "Fácil",
        "salida": datetime(2026, 8, 1, 6, 0),
        "regreso": datetime(2026, 8, 1, 18, 0),
        "placa": "TBQ-UNI-01",
    },
    {
        "hoja": "DUNAS (01-08)",
        "cobranza": "COBRANZA DUNAS (01-08)",
        "destino": "Médanos de Coro (Dunas)",
        "descripcion": "Full day a los Médanos de Coro, estado Falcón.",
        "precio": Decimal("30.00"),
        "dificultad": "Fácil",
        "salida": datetime(2026, 8, 1, 6, 0),
        "regreso": datetime(2026, 8, 1, 18, 0),
        "placa": "TBQ-UNI-02",
    },
    {
        "hoja": "CUBIRO (02-08)",
        "cobranza": "COBRANZA CUBIRO (02-08)",
        "destino": "Cubiro",
        "descripcion": "Paseo de montaña a Cubiro, estado Lara.",
        "precio": Decimal("16.00"),
        "dificultad": "Moderado",
        "salida": datetime(2026, 8, 2, 6, 0),
        "regreso": datetime(2026, 8, 2, 18, 0),
        "placa": "TBQ-UNI-01",
    },
    {
        "hoja": "VARADERO (08-08)",
        "cobranza": "COBRANZA VARADERO (08-08)",
        "destino": "Cayo Sal y Varadero",
        "descripcion": "Full day combinado a Cayo Sal y Varadero, costa de Falcón.",
        "precio": Decimal("25.00"),
        "dificultad": "Fácil",
        "salida": datetime(2026, 8, 8, 6, 0),
        "regreso": datetime(2026, 8, 8, 18, 0),
        "placa": "TBQ-UNI-01",
    },
    {
        "hoja": "CAYO MUERTO (15-08)",
        "cobranza": "COBRANZA CAYO MUERTO (15-08)",
        "destino": "Cayo Muerto",
        "descripcion": "Full day a Cayo Muerto, Parque Nacional Morrocoy. Salida desde Barquisimeto.",
        "precio": Decimal("23.00"),
        "dificultad": "Fácil",
        "salida": datetime(2026, 8, 15, 6, 0),
        "regreso": datetime(2026, 8, 15, 18, 0),
        "placa": "TBQ-UNI-01",
    },
    {
        "hoja": "CAYO MUERTO 2 (15-08)",
        "cobranza": "COBRANZA CAYO MUERTO 2 (15-08)",
        "destino": "Cayo Muerto",
        "descripcion": "Full day a Cayo Muerto, Parque Nacional Morrocoy. Salida desde Barquisimeto.",
        "precio": Decimal("23.00"),
        "dificultad": "Fácil",
        "salida": datetime(2026, 8, 15, 6, 0),
        "regreso": datetime(2026, 8, 15, 18, 0),
        "placa": "TBQ-UNI-02",
    },
    {
        "hoja": "TOUR FALCON (15-08)",
        "cobranza": "COBRANZA TOUR FALCON (15-08)",
        "destino": "Tour Falcón",
        "descripcion": "Tour de dos días por el estado Falcón (sábado y domingo).",
        "precio": Decimal("75.00"),
        "dificultad": "Moderado",
        "salida": datetime(2026, 8, 15, 6, 0),
        "regreso": datetime(2026, 8, 16, 18, 0),
        "placa": "TBQ-UNI-03",
    },
    {
        "hoja": "CASCADA DEL VINO (16-08)",
        "cobranza": "COBRANZA CASCADA VINO (16-08)",
        "destino": "Cascada del Vino",
        "descripcion": "Full day a la Cascada del Vino, estado Lara.",
        "precio": Decimal("23.00"),
        "dificultad": "Moderado",
        "salida": datetime(2026, 8, 16, 6, 0),
        "regreso": datetime(2026, 8, 16, 18, 0),
        "placa": "TBQ-UNI-01",
    },
    {
        "hoja": "PLAYA (22-08)",
        "cobranza": "COBRANZA PLAYA (22-08)",
        "destino": "Cayo Muerto y Tour Chichiriviche",
        "descripcion": "Full day combinado: Cayo Muerto y tour por Chichiriviche, Morrocoy.",
        "precio": Decimal("23.00"),
        "dificultad": "Fácil",
        "salida": datetime(2026, 8, 22, 6, 0),
        "regreso": datetime(2026, 8, 22, 18, 0),
        "placa": "TBQ-UNI-01",
    },
    {
        "hoja": "LA GRIETA (22-08)",
        "cobranza": "COBRANZA GRIETA (22-08)",
        "destino": "La Grieta",
        "descripcion": "Full day a La Grieta, estado Lara.",
        "precio": Decimal("30.00"),
        "dificultad": "Moderado",
        "salida": datetime(2026, 8, 22, 6, 0),
        "regreso": datetime(2026, 8, 22, 18, 0),
        "placa": "TBQ-UNI-02",
    },
    {
        "hoja": "CHORONI (22-08)",
        "cobranza": "COBRANZA CHORONI (22-08)",
        "destino": "Choroní",
        "descripcion": "Tour de dos días a Choroní, estado Aragua.",
        "precio": Decimal("85.00"),
        "dificultad": "Moderado",
        "salida": datetime(2026, 8, 22, 6, 0),
        "regreso": datetime(2026, 8, 23, 18, 0),
        "placa": "TBQ-UNI-03",
    },
    {
        "hoja": "CAYO SOMBRERO (23-08)",
        "cobranza": "COBRANZA CAYO SOMBRERO (23-08)",
        "destino": "Cayo Sombrero y Juanes",
        "descripcion": "Full day a Cayo Sombrero y Cayo Juanes, Parque Nacional Morrocoy.",
        "precio": Decimal("31.00"),
        "dificultad": "Fácil",
        "salida": datetime(2026, 8, 23, 6, 0),
        "regreso": datetime(2026, 8, 23, 18, 0),
        "placa": "TBQ-UNI-01",
    },
]

EMPRESAS = {"NAGUARA TOURS", "TOUR AZUL"}
RE_PUESTO = re.compile(r"^A-\d{2}$")
RE_DIGITO_REF = re.compile(r"^(\d{3,10})\s+(.+)$")


def texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def vacio(valor) -> bool:
    if valor is None:
        return True
    if isinstance(valor, (int, float, Decimal)) and valor == 0:
        return True
    return texto(valor) in {"", "0", "N/E", "N/T", "LIBRE"}


def a_decimal(valor) -> Decimal | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, Decimal):
        return valor
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    bruto = texto(valor).replace(",", ".")
    if not bruto:
        return None
    try:
        return Decimal(bruto)
    except InvalidOperation:
        return None


def titulo_es(valor: str) -> str:
    limpio = re.sub(r"\s+", " ", valor).strip()
    if not limpio:
        return ""
    return " ".join(p[:1].upper() + p[1:].lower() if p else p for p in limpio.split(" "))


def partir_nombre(nombre_crudo: str) -> tuple[str, str, bool, str | None]:
    original = re.sub(r"\s+", " ", nombre_crudo).strip()
    es_menor = bool(re.search(r"NIÑ[OA]|NINO|NINA", original, re.I))
    razon = None
    clave = original.upper()
    if clave in EMPRESAS:
        partes = titulo_es(original).split(" ", 1)
        nombre = partes[0]
        apellido = partes[1] if len(partes) > 1 else "C.A."
        return nombre[:80], apellido[:80], False, titulo_es(original)[:160]

    limpio = re.sub(r"\s*\+\s*NIÑ[OA].*$", "", original, flags=re.I)
    limpio = limpio.replace("&", " ")
    limpio = re.sub(r"\s+", " ", limpio).strip()
    partes = [titulo_es(p) for p in limpio.split(" ") if p]
    if not partes:
        return "Pasajero", "S/N", es_menor, razon
    if len(partes) == 1:
        return partes[0][:80], "S/N", es_menor, razon
    return partes[0][:80], " ".join(partes[1:])[:80], es_menor, razon


def cedula_digitos(valor) -> str | None:
    if vacio(valor):
        return None
    bruto = texto(valor).upper()
    if bruto in {"NIÑO", "NINO", "NIÑA", "N/E", "N/T"}:
        return None
    digitos = re.sub(r"\D", "", bruto)
    if 4 <= len(digitos) <= 9:
        return digitos.lstrip("0") or digitos
    return None


def telefono_limpio(valor) -> str | None:
    if vacio(valor):
        return None
    limpio = re.sub(r"[\s\-().]", "", texto(valor))
    if re.search(r"[A-Za-zÁÉÍÓÚñ]", limpio):
        return None
    if re.fullmatch(r"\+?\d{7,20}", limpio):
        return limpio
    return None


def es_puesto(valor) -> str | None:
    bruto = texto(valor).upper()
    return bruto if RE_PUESTO.fullmatch(bruto) else None


def fila_vacia(row) -> bool:
    return all(v in (None, "") for v in row)


def parsear_guia(texto_guia: str) -> list[tuple[str, str]]:
    crudo = re.sub(r"(?i)^gu[ií]a asignado:\s*", "", texto_guia).strip()
    if not crudo:
        return []
    partes = re.split(r"\s*(?:&| Y )\s*", crudo, flags=re.I)
    guias = []
    for parte in partes:
        nombre, apellido, _, _ = partir_nombre(parte)
        if nombre:
            guias.append((nombre, apellido))
    return guias


def correo_guia(nombre: str, apellido: str) -> str:
    base = f"{nombre}.{apellido}".lower()
    base = (
        base.replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ñ", "n")
        .replace(" ", "")
        .replace("/", "")
    )
    base = re.sub(r"[^a-z0-9.]", "", base).strip(".")
    if base.endswith(".sn"):
        base = base[:-3]
    return f"{base or 'guia'}@travelbqto.com"


def leer_hoja_pasajeros(ws) -> tuple[str, list[dict]]:
    guia = ""
    pasajeros = []
    encabezado = False
    for i, row in enumerate(ws.iter_rows(min_row=1, max_col=15, values_only=True), 1):
        textos = [texto(v).upper() for v in row]
        if i == 2:
            guia = " ".join(texto(v) for v in row if not vacio(v))
        if "PUESTO" in textos and any("NOMBRE" in t for t in textos):
            encabezado = True
            continue
        if not encabezado:
            continue
        puesto = None
        for v in row:
            puesto = es_puesto(v)
            if puesto:
                break
        if not puesto:
            continue
        cols = list(row)
        nombre = cols[7] if len(cols) > 7 else None
        if vacio(nombre) or texto(nombre).upper() == "LIBRE":
            continue
        pasajeros.append(
            {
                "puesto": puesto,
                "nombre": texto(nombre),
                "cedula": cedula_digitos(cols[8] if len(cols) > 8 else None),
                "telefono": telefono_limpio(cols[9] if len(cols) > 9 else None),
                "direccion": None if vacio(cols[10] if len(cols) > 10 else None) else texto(cols[10])[:255],
                "obs": None if vacio(cols[11] if len(cols) > 11 else None) else texto(cols[11])[:255],
            }
        )
    return guia, pasajeros


def leer_resumen_cobranza(ws) -> dict[str, dict]:
    resumen = {}
    seccion = False
    for row in ws.iter_rows(min_row=1, max_col=8, values_only=True):
        textos = [texto(v).upper() for v in row]
        if row[0] and texto(row[0]).upper() == "DETALLES DE PAGOS":
            break
        if "PUESTO" in textos and any("NOMBRE" in t for t in textos) and any("MONTO" in t for t in textos):
            seccion = True
            continue
        if not seccion:
            continue
        puesto = es_puesto(row[0])
        if not puesto:
            continue
        if vacio(row[1]) or texto(row[1]).upper() == "LIBRE":
            continue
        resumen[puesto] = {
            "nombre": texto(row[1]),
            "monto": a_decimal(row[2]) or Decimal("0"),
            "abono": a_decimal(row[3]) or Decimal("0"),
            "saldo": a_decimal(row[4]) or Decimal("0"),
            "obs": None if vacio(row[5]) else texto(row[5]),
        }
    return resumen


def extraer_pago(row) -> dict | None:
    abono = row[2] if not isinstance(row[2], str) else None
    notas_abono = texto(row[2]) if isinstance(row[2], str) and not vacio(row[2]) else ""
    tasa = a_decimal(row[3])
    usd = a_decimal(row[4])
    referencia = None if vacio(row[5]) else texto(row[5])
    fecha = row[6]
    extra = " | ".join(texto(v) for v in row[7:] if not vacio(v))
    notas = " | ".join(p for p in (notas_abono, extra) if p)
    if abono is None:
        abono_dec = None
    else:
        abono_dec = a_decimal(abono)
    if abono_dec is None and usd is None and not referencia and not notas:
        return None
    fecha_pago = None
    if isinstance(fecha, datetime):
        fecha_pago = fecha.date()
    elif isinstance(fecha, date):
        fecha_pago = fecha
    return {
        "abono_bs": abono_dec,
        "tasa": tasa,
        "usd": usd,
        "referencia": referencia,
        "fecha": fecha_pago,
        "notas": notas[:255] if notas else None,
    }


def clasificar_pago(pago: dict) -> tuple[str, str | None, str]:
    ref = (pago.get("referencia") or "").upper()
    notas = (pago.get("notas") or "").upper()
    combinado = f"{ref} {notas}"
    if "EFECTIVO" in combinado:
        return "efectivo_usd", None, pago.get("referencia") or "EFECTIVO"
    if "BINANCE" in combinado or "USDT" in combinado:
        return "zelle", None, pago.get("referencia") or "BINANCE"

    banco_codigo = None
    referencia = pago.get("referencia") or ""
    if referencia:
        m = RE_DIGITO_REF.match(referencia)
        if m:
            referencia, banco_txt = m.group(1), m.group(2)
        else:
            banco_txt = referencia
        for alias, codigo in ALIAS_BANCO.items():
            if re.search(rf"\b{alias}\b", banco_txt.upper()):
                banco_codigo = codigo
                break
        if banco_codigo and m:
            pass
        elif banco_codigo and not m:
            # "PROVINCIAL MARIA" / "PROVINCIAL"
            if banco_txt.upper().strip() in ALIAS_BANCO:
                referencia = referencia or banco_txt
    metodo = "transferencia"
    if pago.get("abono_bs") is None and pago.get("usd") is not None:
        if banco_codigo:
            metodo = "transferencia"
        else:
            metodo = "zelle" if "BINANCE" in combinado else "efectivo_usd"
    return metodo, banco_codigo, (referencia or "")[:120]


def leer_grupos_pago(ws) -> list[dict]:
    grupos: list[dict] = []
    actual: dict | None = None
    pendientes: list[dict] = []
    seccion = False
    encabezado = False

    def cerrar():
        nonlocal actual
        if actual and (actual["puestos"] or actual["pagos"]):
            grupos.append(actual)
        actual = None

    def abrir():
        nonlocal actual
        actual = {"puestos": [], "pagos": list(pendientes), "pagador": None}
        pendientes.clear()

    for row in ws.iter_rows(min_row=1, max_col=10, values_only=True):
        if row[0] and texto(row[0]).upper() == "DETALLES DE PAGOS":
            seccion = True
            continue
        if not seccion:
            continue
        textos = [texto(v).upper() for v in row]
        if not encabezado and "PUESTO" in textos and any("NOMBRE" in t for t in textos):
            encabezado = True
            continue
        if not encabezado:
            continue
        if fila_vacia(row):
            cerrar()
            continue
        puesto = es_puesto(row[0])
        nombre = None if vacio(row[1]) else texto(row[1])
        if nombre and nombre.upper() == "LIBRE":
            continue
        pago = extraer_pago(row)
        if puesto:
            if actual is None:
                abrir()
            actual["puestos"].append(puesto)
            if nombre and not actual["pagador"]:
                actual["pagador"] = nombre
            if pago:
                actual["pagos"].append(pago)
        elif pago:
            if actual is None:
                pendientes.append(pago)
            else:
                actual["pagos"].append(pago)
    cerrar()
    return grupos


def conectar():
    return pymysql.connect(
        host=os.getenv("DB_HOST", "127.0.0.1"),
        port=int(os.getenv("DB_PUERTO", "3306")),
        user=os.getenv("DB_USUARIO", "root"),
        password=os.getenv("DB_CONTRASENA", ""),
        database=os.getenv("DB_NOMBRE", "travel_bqto"),
        charset="utf8mb4",
        autocommit=False,
        cursorclass=pymysql.cursors.DictCursor,
    )


def fetchone(cur, sql, args=None):
    cur.execute(sql, args or ())
    return cur.fetchone()


def fetchall(cur, sql, args=None):
    cur.execute(sql, args or ())
    return cur.fetchall()


def insertar(cur, sql, args) -> int:
    cur.execute(sql, args)
    return cur.lastrowid


def asegurar_catalogo(cur, ahora: datetime) -> dict:
    cur.execute("SET @sigel_usuario_id = %s", (ADMIN_ID,))
    for codigo, nombre in BANCOS_NUEVOS:
        existe = fetchone(cur, "SELECT id FROM bancos WHERE codigo=%s", (codigo,))
        if not existe:
            insertar(
                cur,
                "INSERT INTO bancos (codigo, nombre, activo, creado_en, actualizado_en) "
                "VALUES (%s,%s,1,%s,%s)",
                (codigo, nombre, ahora, ahora),
            )
    bancos = {r["codigo"]: r["id"] for r in fetchall(cur, "SELECT id, codigo FROM bancos")}
    metodos = {r["codigo"]: r["id"] for r in fetchall(cur, "SELECT id, codigo FROM metodos_pago")}
    unidades = {}
    asientos = {}
    for placa, modelo, capacidad in UNIDADES:
        fila = fetchone(cur, "SELECT id FROM unidades_transporte WHERE placa=%s AND eliminado_en IS NULL", (placa,))
        if fila:
            unidad_id = fila["id"]
        else:
            unidad_id = insertar(
                cur,
                "INSERT INTO unidades_transporte (placa, modelo, capacidad, creado_en, actualizado_en) "
                "VALUES (%s,%s,%s,%s,%s)",
                (placa, modelo, capacidad, ahora, ahora),
            )
        unidades[placa] = unidad_id
        existentes = {
            r["numero"]: r["id"]
            for r in fetchall(
                cur,
                "SELECT id, numero FROM asientos WHERE unidad_id=%s AND eliminado_en IS NULL",
                (unidad_id,),
            )
        }
        mapa = {}
        for n in range(1, 32):
            numero = f"A-{n:02d}"
            if numero in existentes:
                mapa[numero] = existentes[numero]
                continue
            posicion = "ventana" if n % 2 == 1 else "pasillo"
            asiento_id = insertar(
                cur,
                "INSERT INTO asientos (unidad_id, numero, posicion, creado_en, actualizado_en) "
                "VALUES (%s,%s,%s,%s,%s)",
                (unidad_id, numero, posicion, ahora, ahora),
            )
            mapa[numero] = asiento_id
        asientos[placa] = mapa
    return {"bancos": bancos, "metodos": metodos, "unidades": unidades, "asientos": asientos}


def asegurar_destino(cur, cfg, ahora: datetime) -> int:
    fila = fetchone(cur, "SELECT id FROM destinos WHERE nombre=%s AND eliminado_en IS NULL", (cfg["destino"],))
    if fila:
        cur.execute(
            "UPDATE destinos SET descripcion=%s, precio_base_eur=%s, dificultad=%s, "
            "activo=1, actualizado_en=%s WHERE id=%s",
            (cfg["descripcion"], cfg["precio"], cfg["dificultad"], ahora, fila["id"]),
        )
        return fila["id"]
    return insertar(
        cur,
        "INSERT INTO destinos (nombre, descripcion, precio_base_eur, recargo_menor_eur, "
        "dificultad, activo, creado_en, actualizado_en) VALUES (%s,%s,%s,0,%s,1,%s,%s)",
        (cfg["destino"], cfg["descripcion"], cfg["precio"], cfg["dificultad"], ahora, ahora),
    )


def asegurar_guia(cur, nombre: str, apellido: str, ahora: datetime) -> int:
    correo = correo_guia(nombre, apellido)
    tabla_usuarios = f"`{BD_SEGURIDAD}`.usuarios"
    fila = fetchone(
        cur,
        f"SELECT id FROM {tabla_usuarios} WHERE correo=%s AND eliminado_en IS NULL",
        (correo,),
    )
    if fila:
        return fila["id"]
    return insertar(
        cur,
        f"INSERT INTO {tabla_usuarios} "
        "(rol_id, correo, hash_contrasena, nombre, apellido, creado_en, actualizado_en) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s)",
        (ROL_GUIA, correo, HASH_CLAVE, nombre[:80], apellido[:80], ahora, ahora),
    )


def asegurar_tasa(cur, fecha_pago: date, valor: Decimal, cache: dict, ahora: datetime) -> int:
    clave = (fecha_pago.isoformat(), str(valor))
    if clave in cache:
        return cache[clave]
    fila = fetchone(
        cur,
        "SELECT id FROM tasas WHERE fecha=%s AND valor=%s AND moneda_id=1 AND eliminado_en IS NULL",
        (fecha_pago, valor),
    )
    if fila:
        cache[clave] = fila["id"]
        return fila["id"]
    tasa_id = insertar(
        cur,
        "INSERT INTO tasas (fecha, valor, moneda_id, origen) VALUES (%s,%s,1,'manual')",
        (fecha_pago, valor),
    )
    cache[clave] = tasa_id
    return tasa_id


def asegurar_punto(cur, direccion: str, cache: dict, ahora: datetime) -> int:
    clave = direccion.strip().upper()
    if clave in cache:
        return cache[clave]
    fila = fetchone(
        cur,
        "SELECT id FROM puntos_recogida WHERE UPPER(direccion)=%s AND tipo='domicilio' "
        "AND eliminado_en IS NULL",
        (clave,),
    )
    if fila:
        cache[clave] = fila["id"]
        return fila["id"]
    nombre = direccion[:120]
    punto_id = insertar(
        cur,
        "INSERT INTO puntos_recogida (nombre, direccion, ciudad, estado, tipo, activo, "
        "creado_por, creado_en, actualizado_en) VALUES (%s,%s,'Barquisimeto','Lara','domicilio',1,%s,%s,%s)",
        (nombre, direccion[:255], ADMIN_ID, ahora, ahora),
    )
    cache[clave] = punto_id
    return punto_id


def vincular_punto_cliente(cur, cliente_id: int, punto_id: int, ahora: datetime) -> None:
    existe = fetchone(
        cur,
        "SELECT id FROM clientes_puntos_recogida WHERE cliente_id=%s AND punto_recogida_id=%s "
        "AND eliminado_en IS NULL",
        (cliente_id, punto_id),
    )
    if existe:
        return
    insertar(
        cur,
        "INSERT INTO clientes_puntos_recogida (cliente_id, punto_recogida_id, es_predeterminado, "
        "creado_en, actualizado_en) VALUES (%s,%s,1,%s,%s)",
        (cliente_id, punto_id, ahora, ahora),
    )


class Importador:
    def __init__(self, cur):
        self.cur = cur
        self.ci_seq = CI_INTERNO_INICIO
        self.clientes_ci: dict[str, int] = {}
        self.empresas: dict[str, int] = {}
        for fila in fetchall(
            cur,
            "SELECT id, tipo_documento, numero_documento FROM clientes WHERE eliminado_en IS NULL",
        ):
            if fila["tipo_documento"] == "V":
                self.clientes_ci[fila["numero_documento"]] = fila["id"]
            elif fila["tipo_documento"] == "J":
                self.empresas[fila["numero_documento"]] = fila["id"]

    def siguiente_ci(self) -> str:
        while True:
            numero = str(self.ci_seq)
            self.ci_seq += 1
            if numero not in self.clientes_ci:
                return numero

    def cliente_empresa(self, razon: str, ahora: datetime) -> int:
        if razon in self.empresas:
            return self.empresas[razon]
        numero = f"J{len(self.empresas)+1:07d}"
        nombre, apellido, _, _ = partir_nombre(razon)
        cliente_id = insertar(
            self.cur,
            "INSERT INTO clientes (tipo_cliente, tipo_documento, numero_documento, nombre, apellido, "
            "razon_social, estado_id, ciudad_id, notas, creado_por, actualizado_por, creado_en, actualizado_en) "
            "VALUES ('juridico','J',%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
            (
                numero,
                nombre,
                apellido,
                razon[:160],
                ESTADO_LARA,
                CIUDAD_BQTO,
                MARCA,
                ADMIN_ID,
                ADMIN_ID,
                ahora,
                ahora,
            ),
        )
        self.empresas[razon] = cliente_id
        return cliente_id

    def crear_cliente(self, pax: dict, ahora: datetime, como_empresa: bool = False) -> int:
        nombre, apellido, es_menor, razon = partir_nombre(pax["nombre"])
        if razon and como_empresa:
            return self.cliente_empresa(razon, ahora)
        if razon:
            nombre, apellido = "Pasajero", razon[:80]
            pax = dict(pax)
            pax["cedula"] = None
        ci = pax.get("cedula")
        notas = [MARCA]
        if pax.get("obs"):
            notas.append(pax["obs"])
        if es_menor:
            notas.append("Menor según planilla")
        if ci and ci in self.clientes_ci:
            return self.clientes_ci[ci]
        if not ci:
            ci = self.siguiente_ci()
            notas.append("Cédula no reportada en la planilla; código interno de importación")
        cliente_id = insertar(
            self.cur,
            "INSERT INTO clientes (tipo_cliente, tipo_documento, numero_documento, nombre, apellido, "
            "telefono, direccion, estado_id, ciudad_id, notas, creado_por, actualizado_por, creado_en, actualizado_en) "
            "VALUES ('natural','V',%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
            (
                ci,
                nombre,
                apellido,
                pax.get("telefono"),
                pax.get("direccion"),
                ESTADO_LARA,
                CIUDAD_BQTO,
                " | ".join(notas)[:65535],
                ADMIN_ID,
                ADMIN_ID,
                ahora,
                ahora,
            ),
        )
        self.clientes_ci[ci] = cliente_id
        return cliente_id


def agrupar_asientos(pasajeros: list[dict], grupos_pago: list[dict], resumen: dict) -> list[dict]:
    por_puesto = {p["puesto"]: p for p in pasajeros}
    usados = set()
    reservas = []
    for grupo in grupos_pago:
        puestos = [p for p in grupo["puestos"] if p in por_puesto and p not in usados]
        if not puestos:
            continue
        usados.update(puestos)
        reservas.append({"puestos": puestos, "pagos": [p for p in grupo["pagos"] if p], "pagador": grupo.get("pagador")})

    restantes = [p for p in pasajeros if p["puesto"] not in usados]
    por_nombre: dict[str, list[dict]] = {}
    for pax in restantes:
        por_nombre.setdefault(pax["nombre"].strip().upper(), []).append(pax)
    for paxs in por_nombre.values():
        puestos = [p["puesto"] for p in paxs]
        reservas.append({"puestos": puestos, "pagos": [], "pagador": paxs[0]["nombre"]})
    return reservas


def estado_reserva(puestos: list[str], resumen: dict) -> str:
    saldos = [resumen[p]["saldo"] for p in puestos if p in resumen]
    abonos = [resumen[p]["abono"] for p in puestos if p in resumen]
    if saldos and all(s <= 0 for s in saldos):
        return "abonada"
    if any(a > 0 for a in abonos):
        return "confirmada"
    return "pendiente"


def monto_pago_sistema(pago: dict, metodo: str) -> Decimal | None:
    if metodo in {"transferencia", "pago_movil", "efectivo_bs"}:
        if pago.get("abono_bs") and pago["abono_bs"] > 0:
            return pago["abono_bs"].quantize(Decimal("0.01"))
        if pago.get("usd") and pago.get("tasa"):
            return (pago["usd"] * pago["tasa"]).quantize(Decimal("0.01"))
    if pago.get("usd") and pago["usd"] > 0:
        return pago["usd"].quantize(Decimal("0.01"))
    if pago.get("abono_bs") and pago["abono_bs"] > 0 and pago.get("tasa"):
        return (pago["abono_bs"] / pago["tasa"]).quantize(Decimal("0.01"))
    return None


def borrar_agosto(cur) -> None:
    ids = [
        r["id"]
        for r in fetchall(
            cur,
            "SELECT id FROM viajes WHERE fecha_salida >= '2026-08-01' AND fecha_salida < '2026-09-01' "
            "AND eliminado_en IS NULL",
        )
    ]
    if not ids:
        return
    formato = ",".join(["%s"] * len(ids))
    reserva_ids = [
        r["id"] for r in fetchall(cur, f"SELECT id FROM reservas WHERE viaje_id IN ({formato})", ids)
    ]
    if reserva_ids:
        rf = ",".join(["%s"] * len(reserva_ids))
        rc_ids = [
            r["id"]
            for r in fetchall(cur, f"SELECT id FROM reserva_clientes WHERE reserva_id IN ({rf})", reserva_ids)
        ]
        if rc_ids:
            rcf = ",".join(["%s"] * len(rc_ids))
            cur.execute(f"DELETE FROM abordajes_viaje WHERE reserva_cliente_id IN ({rcf})", rc_ids)
            cur.execute(f"DELETE FROM asientos_reservados WHERE reserva_cliente_id IN ({rcf})", rc_ids)
            cur.execute(f"DELETE FROM viajes_ruta_recogida WHERE reserva_cliente_id IN ({rcf})", rc_ids)
        cur.execute(f"DELETE FROM pagos WHERE reserva_id IN ({rf})", reserva_ids)
        cur.execute(f"DELETE FROM reserva_clientes WHERE reserva_id IN ({rf})", reserva_ids)
        cur.execute(f"DELETE FROM reservas WHERE id IN ({rf})", reserva_ids)
    cur.execute(f"DELETE FROM viajes_guias WHERE viaje_id IN ({formato})", ids)
    cur.execute(f"DELETE FROM costos_operativos WHERE viaje_id IN ({formato})", ids)
    cur.execute(f"DELETE FROM viajes WHERE id IN ({formato})", ids)


def importar(excel: Path, forzar: bool) -> None:
    if not excel.exists():
        raise SystemExit(f"No está el Excel: {excel}")

    wb = load_workbook(str(excel), read_only=True, data_only=True, keep_links=False)
    ahora = datetime.now()
    conn = conectar()
    cur = conn.cursor()
    try:
        existentes = fetchone(
            cur,
            "SELECT COUNT(*) AS n FROM viajes WHERE fecha_salida >= '2026-08-01' "
            "AND fecha_salida < '2026-09-01' AND eliminado_en IS NULL",
        )["n"]
        if existentes and not forzar:
            raise SystemExit(
                "Ya hay viajes de agosto 2026. Si quieres recargarlos: python importar_planilla_agosto.py --forzar"
            )
        if forzar:
            borrar_agosto(cur)

        cat = asegurar_catalogo(cur, ahora)
        importador = Importador(cur)
        tasas_cache: dict = {}
        puntos_cache: dict = {}
        contadores = {
            "viajes": 0,
            "reservas": 0,
            "pasajeros": 0,
            "pagos": 0,
            "clientes_nuevos": 0,
        }
        clientes_antes = fetchone(cur, "SELECT COUNT(*) n FROM clientes")["n"]

        for cfg in VIAJES:
            guia_txt, pasajeros = leer_hoja_pasajeros(wb[cfg["hoja"]])
            resumen = leer_resumen_cobranza(wb[cfg["cobranza"]])
            grupos = leer_grupos_pago(wb[cfg["cobranza"]])
            destino_id = asegurar_destino(cur, cfg, ahora)
            unidad_id = cat["unidades"][cfg["placa"]]
            viaje = fetchone(
                cur,
                "SELECT id FROM viajes WHERE destino_id=%s AND unidad_id=%s AND fecha_salida=%s "
                "AND eliminado_en IS NULL",
                (destino_id, unidad_id, cfg["salida"]),
            )
            if viaje:
                viaje_id = viaje["id"]
            else:
                viaje_id = insertar(
                    cur,
                    "INSERT INTO viajes (destino_id, unidad_id, fecha_salida, fecha_regreso, estado, "
                    "creado_en, actualizado_en) VALUES (%s,%s,%s,%s,'finalizado',%s,%s)",
                    (destino_id, unidad_id, cfg["salida"], cfg["regreso"], ahora, ahora),
                )
                contadores["viajes"] += 1

            for i, (nombre, apellido) in enumerate(parsear_guia(guia_txt)):
                guia_id = asegurar_guia(cur, nombre, apellido, ahora)
                existe = fetchone(
                    cur,
                    "SELECT id FROM viajes_guias WHERE viaje_id=%s AND usuario_id=%s AND eliminado_en IS NULL",
                    (viaje_id, guia_id),
                )
                if not existe:
                    insertar(
                        cur,
                        "INSERT INTO viajes_guias (viaje_id, usuario_id, es_principal, creado_en, actualizado_en) "
                        "VALUES (%s,%s,%s,%s,%s)",
                        (viaje_id, guia_id, 1 if i == 0 else 0, ahora, ahora),
                    )

            mapa_asientos = cat["asientos"][cfg["placa"]]
            for grupo in agrupar_asientos(pasajeros, grupos, resumen):
                paxs = [next(p for p in pasajeros if p["puesto"] == puesto) for puesto in grupo["puestos"]]
                titular = next((p for p in paxs if p.get("cedula")), paxs[0])
                cliente_titular = importador.crear_cliente(
                    titular, ahora, como_empresa=partir_nombre(titular["nombre"])[3] is not None
                )
                fechas_pago = [p["fecha"] for p in grupo["pagos"] if p.get("fecha")]
                fecha_reserva = min(fechas_pago) if fechas_pago else cfg["salida"].date()
                fecha_reserva_dt = datetime.combine(fecha_reserva, datetime.min.time().replace(hour=10))
                estado = estado_reserva(grupo["puestos"], resumen)
                reserva_id = insertar(
                    cur,
                    "INSERT INTO reservas (cliente_id, viaje_id, fecha_reserva, estado, creado_por, "
                    "creado_en, actualizado_en) VALUES (%s,%s,%s,%s,%s,%s,%s)",
                    (cliente_titular, viaje_id, fecha_reserva_dt, estado, ADMIN_ID, ahora, ahora),
                )
                contadores["reservas"] += 1
                ids_usados = {
                    fetchone(cur, "SELECT numero_documento n FROM clientes WHERE id=%s", (cliente_titular,))["n"]
                }
                puesto_titular = titular["puesto"]
                for pax in paxs:
                    if pax["puesto"] == puesto_titular:
                        cliente_id = cliente_titular
                    else:
                        # Un cliente no puede repetirse en la misma reserva.
                        pax_copia = dict(pax)
                        if pax_copia.get("cedula") and pax_copia["cedula"] in ids_usados:
                            pax_copia["cedula"] = None
                        cliente_id = importador.crear_cliente(pax_copia, ahora, como_empresa=False)
                    ids_usados.add(
                        fetchone(cur, "SELECT numero_documento n FROM clientes WHERE id=%s", (cliente_id,))["n"]
                    )
                    cob = resumen.get(pax["puesto"], {})
                    precio = cob.get("monto") or cfg["precio"]
                    es_menor = bool(re.search(r"NIÑ[OA]|NINO|NINA|EXONERADO", f"{pax['nombre']} {cob.get('obs') or ''}", re.I))
                    punto_id = None
                    if pax.get("direccion"):
                        punto_id = asegurar_punto(cur, pax["direccion"], puntos_cache, ahora)
                        vincular_punto_cliente(cur, cliente_id, punto_id, ahora)
                    rc_id = insertar(
                        cur,
                        "INSERT INTO reserva_clientes (reserva_id, cliente_id, es_titular, es_menor, ocupa_asiento, "
                        "precio_pasajero_eur, recargo_eur, notas_tarifa, punto_recogida_id, creado_en, actualizado_en) "
                        "VALUES (%s,%s,%s,%s,1,%s,0,%s,%s,%s,%s)",
                        (
                            reserva_id,
                            cliente_id,
                            1 if pax["puesto"] == puesto_titular else 0,
                            1 if es_menor else 0,
                            precio,
                            (cob.get("obs") or pax.get("obs") or "")[:255] or None,
                            punto_id,
                            ahora,
                            ahora,
                        ),
                    )
                    insertar(
                        cur,
                        "INSERT INTO asientos_reservados (reserva_cliente_id, viaje_id, asiento_id, creado_en, actualizado_en) "
                        "VALUES (%s,%s,%s,%s,%s)",
                        (rc_id, viaje_id, mapa_asientos[pax["puesto"]], ahora, ahora),
                    )
                    contadores["pasajeros"] += 1

                pagos_grupo = [p for p in grupo["pagos"] if p]
                if not pagos_grupo:
                    total = sum((resumen[p]["abono"] for p in grupo["puestos"] if p in resumen), Decimal("0"))
                    if total > 0:
                        pagos_grupo = [
                            {
                                "abono_bs": None,
                                "tasa": None,
                                "usd": total,
                                "referencia": "PLANILLA AGOSTO",
                                "fecha": fecha_reserva,
                                "notas": "Liquidado en cobranza de la planilla, sin referencia bancaria",
                            }
                        ]
                for pago in pagos_grupo:
                    metodo, banco_codigo, referencia = clasificar_pago(pago)
                    monto = monto_pago_sistema(pago, metodo)
                    if not monto or monto <= 0:
                        continue
                    tasa_valor = pago.get("tasa") or Decimal("160.0000")
                    fecha_pago = pago.get("fecha") or fecha_reserva
                    tasa_id = asegurar_tasa(cur, fecha_pago, tasa_valor, tasas_cache, ahora)
                    tipo = "total" if len(pagos_grupo) == 1 and estado == "abonada" else "cuota"
                    notas = pago.get("notas")
                    insertar(
                        cur,
                        "INSERT INTO pagos (reserva_id, metodo_pago_id, tasa_id, monto, tipo, estado, fecha_pago, "
                        "referencia, banco_origen_id, validado_por, validado_en, notas, creado_por, creado_en, actualizado_en) "
                        "VALUES (%s,%s,%s,%s,%s,'aprobado',%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                        (
                            reserva_id,
                            cat["metodos"][metodo],
                            tasa_id,
                            monto,
                            tipo,
                            fecha_pago,
                            referencia or None,
                            cat["bancos"].get(banco_codigo) if banco_codigo else None,
                            ADMIN_ID,
                            ahora,
                            notas,
                            ADMIN_ID,
                            ahora,
                            ahora,
                        ),
                    )
                    contadores["pagos"] += 1

        clientes_despues = fetchone(cur, "SELECT COUNT(*) n FROM clientes")["n"]
        contadores["clientes_nuevos"] = clientes_despues - clientes_antes
        conn.commit()
        print("Planilla de agosto 2026 cargada.")
        for k, v in contadores.items():
            print(f"  {k}: {v}")
    except Exception:
        conn.rollback()
        raise
    finally:
        wb.close()
        cur.close()
        conn.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Importa la planilla real de agosto 2026 a SIGEL")
    parser.add_argument("--excel", type=Path, default=EXCEL_DEFECTO)
    parser.add_argument("--forzar", action="store_true", help="Borra viajes de agosto 2026 y vuelve a cargar")
    args = parser.parse_args()
    importar(args.excel, args.forzar)


if __name__ == "__main__":
    main()
